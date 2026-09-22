"""Render and persist Phase 3 quotation exports from immutable snapshots."""

from __future__ import annotations

import hashlib
import re
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as ReportLabImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.boq_v2 import (
    BOQV2CostPlan,
    BOQV2CostSelection,
    BOQV2ExportArtifact,
    BOQV2RevisionSnapshot,
    BOQV2Vendor,
    BOQV2VendorOffer,
    BOQV2VendorOfferLine,
)
from app.schemas.boq_quotation_schema import (
    BOQV2ExportArtifactResponse,
    BOQV2ExportRequest,
)
from app.services.boq_document_service import (
    _begin_command,
    _complete_command,
    _revision_context,
)
from app.services.boq_domain_service import BOQDomainError
from app.services.gcs_storage_service import (
    download_storage_key_bytes,
    generate_signed_url_for_storage_key,
    upload_boq_export_artifact,
)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"
FORMULA_PREFIXES = ("=", "+", "-", "@")
THIN_GREY = Side(style="thin", color="D9E0E7")
HEADER_FILL = PatternFill("solid", fgColor="15324A")
SUBHEADER_FILL = PatternFill("solid", fgColor="E7EEF4")
TOTAL_FILL = PatternFill("solid", fgColor="EAF2EE")


def _now() -> datetime:
    return datetime.now(UTC)


def _timestamp(value: datetime | None) -> str:
    current = value or _now()
    if current.tzinfo is None:
        current = current.replace(tzinfo=UTC)
    return current.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _safe_text(value: object | None) -> str:
    """Keep spreadsheet text inert when opened by office applications."""

    text = str(value or "")
    if text.lstrip().startswith(FORMULA_PREFIXES):
        return f"'{text}"
    return text


def _decimal(value: object | None) -> Decimal:
    return Decimal(str(value or "0"))


def _number(value: object | None) -> float:
    return float(_decimal(value))


def _set_text(cell: Cell, value: object | None) -> None:
    cell.value = _safe_text(value)
    cell.data_type = "s"


def _payload(snapshot: BOQV2RevisionSnapshot, audience: str) -> dict[str, Any]:
    raw = snapshot.customer_payload if audience == "CUSTOMER" else snapshot.internal_payload
    if not isinstance(raw, dict):
        raise BOQDomainError(
            "EXPORT_SNAPSHOT_INVALID", "Quotation snapshot payload is invalid"
        )
    return dict(raw)


def _included_scope(payload: dict[str, Any]) -> list[dict[str, Any]]:
    scope = payload.get("scope") or []
    return [
        dict(node)
        for node in scope
        if isinstance(node, dict) and node.get("inclusion_state") != "EXCLUDED"
    ]


def _style_worksheet(sheet: Any, *, internal: bool) -> None:
    widths = [7, 14, 38, 34, 12, 12, 16, 16, 17]
    if internal:
        widths.extend([17, 17])
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A12"
    sheet.print_title_rows = "11:11"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if internal else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = True
    sheet.oddFooter.center.text = "Page &P of &N"


def _write_workbook(payload: dict[str, Any], *, internal: bool) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Internal BOQ" if internal else "Quotation"
    _style_worksheet(sheet, internal=internal)

    document = dict(payload.get("document") or {})
    revision = dict(payload.get("revision") or {})
    project = dict(payload.get("project") or {})
    quotation = dict(payload.get("quotation") or {})
    calculation = dict(payload.get("calculation") or {})
    max_column = 11 if internal else 9

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_column)
    _set_text(sheet.cell(1, 1), quotation.get("title") or "ใบเสนอราคา / Quotation")
    sheet.cell(1, 1).font = Font(size=18, bold=True, color="15324A")
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")

    header_rows = [
        ("เลขที่ / No.", document.get("number"), "Revision", revision.get("number")),
        ("โครงการ / Project", project.get("name"), "วันที่ / Date", quotation.get("quotation_date")),
        ("ลูกค้า / Customer", quotation.get("customer_name"), "ใช้ได้ถึง / Valid until", quotation.get("valid_until")),
        ("ที่อยู่ / Address", quotation.get("customer_address"), "เลขประจำตัวผู้เสียภาษี", quotation.get("customer_tax_id")),
        ("ผู้ติดต่อ / Contact", quotation.get("customer_contact"), "สกุลเงิน / Currency", quotation.get("currency") or "THB"),
    ]
    for row_index, values in enumerate(header_rows, start=3):
        _set_text(sheet.cell(row_index, 1), values[0])
        _set_text(sheet.cell(row_index, 2), values[1])
        _set_text(sheet.cell(row_index, 6), values[2])
        _set_text(sheet.cell(row_index, 7), values[3])
        sheet.cell(row_index, 1).font = Font(bold=True)
        sheet.cell(row_index, 6).font = Font(bold=True)
        sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=5)
        sheet.merge_cells(start_row=row_index, start_column=7, end_row=row_index, end_column=max_column)

    columns = [
        "ลำดับ",
        "รหัส",
        "รายละเอียด",
        "ข้อกำหนด",
        "จำนวน",
        "หน่วย",
        "ค่าวัสดุ/หน่วย",
        "ค่าแรง/หน่วย",
        "รวม",
    ]
    if internal:
        columns.extend(["ต้นทุนรวม", "กำไรขั้นต้น"])
    for column_index, label in enumerate(columns, start=1):
        cell = sheet.cell(11, column_index)
        _set_text(cell, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)
    sheet.row_dimensions[11].height = 32

    row_index = 12
    item_number = 0
    for node in _included_scope(payload):
        kind = str(node.get("node_kind") or "")
        if kind != "ITEM":
            sheet.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=max_column,
            )
            _set_text(sheet.cell(row_index, 1), node.get("description") or kind.title())
            sheet.cell(row_index, 1).font = Font(bold=True, color="15324A")
            sheet.cell(row_index, 1).fill = SUBHEADER_FILL
            sheet.cell(row_index, 1).alignment = Alignment(
                indent=min(int(node.get("depth") or 0), 6), wrap_text=True
            )
            row_index += 1
            continue

        item_number += 1
        sheet.cell(row_index, 1, item_number)
        _set_text(sheet.cell(row_index, 2), node.get("item_code"))
        _set_text(sheet.cell(row_index, 3), node.get("description"))
        _set_text(sheet.cell(row_index, 4), node.get("specification"))
        sheet.cell(row_index, 5, _number(node.get("quantity")))
        _set_text(sheet.cell(row_index, 6), node.get("unit"))
        sheet.cell(row_index, 7, _number(node.get("sell_material_unit_rate")))
        sheet.cell(row_index, 8, _number(node.get("sell_labor_unit_rate")))
        sheet.cell(row_index, 9, _number(node.get("sell_total")))
        if internal:
            cost_total = sum(
                _decimal(component.get("total"))
                for component in node.get("components") or []
                if isinstance(component, dict) and component.get("total") is not None
            )
            sheet.cell(row_index, 10, float(cost_total))
            sheet.cell(
                row_index,
                11,
                float(_decimal(node.get("sell_total")) - cost_total),
            )
        for column_index in range(1, max_column + 1):
            cell = sheet.cell(row_index, column_index)
            cell.border = Border(bottom=THIN_GREY)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index in {2, 3, 4, 6},
                horizontal="right" if column_index in {5, 7, 8, 9, 10, 11} else "left",
            )
        for column_index in range(5, max_column + 1):
            sheet.cell(row_index, column_index).number_format = "#,##0.00"
        row_index += 1

    total_start = row_index + 1
    totals = [
        ("รวมก่อนส่วนลด / Subtotal", calculation.get("subtotal")),
        ("ส่วนลด / Discount", calculation.get("discount_amount")),
        ("รวมก่อน VAT / Net", calculation.get("net_sell_ex_vat")),
        (f"VAT {calculation.get('vat_rate') or '0'}%", calculation.get("vat_amount")),
        ("ยอดรวมสุทธิ / Grand total", calculation.get("grand_total")),
    ]
    for offset, (label, value) in enumerate(totals):
        current = total_start + offset
        sheet.merge_cells(
            start_row=current,
            start_column=1,
            end_row=current,
            end_column=max_column - 1,
        )
        _set_text(sheet.cell(current, 1), label)
        sheet.cell(current, 1).alignment = Alignment(horizontal="right")
        sheet.cell(current, 1).font = Font(bold=offset == len(totals) - 1)
        sheet.cell(current, max_column, _number(value))
        sheet.cell(current, max_column).number_format = '#,##0.00 "THB"'
        if offset == len(totals) - 1:
            for column_index in range(1, max_column + 1):
                sheet.cell(current, column_index).fill = TOTAL_FILL
                sheet.cell(current, column_index).font = Font(bold=True)

    row_index = total_start + len(totals) + 2
    _set_text(sheet.cell(row_index, 1), "เงื่อนไขการชำระเงิน / Payment schedule")
    sheet.cell(row_index, 1).font = Font(bold=True, color="15324A")
    row_index += 1
    for schedule in quotation.get("payment_schedule") or []:
        if not isinstance(schedule, dict):
            continue
        _set_text(sheet.cell(row_index, 1), f"• {schedule.get('label') or ''}")
        sheet.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=max_column - 1,
        )
        sheet.cell(row_index, max_column, _number(schedule.get("amount")))
        sheet.cell(row_index, max_column).number_format = '#,##0.00 "THB"'
        row_index += 1

    row_index += 1
    _set_text(sheet.cell(row_index, 1), "เงื่อนไขการค้า / Commercial terms")
    sheet.cell(row_index, 1).font = Font(bold=True, color="15324A")
    row_index += 1
    for term in quotation.get("commercial_terms") or []:
        _set_text(sheet.cell(row_index, 1), f"• {term}")
        sheet.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=max_column,
        )
        sheet.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="top")
        row_index += 1

    sheet.auto_filter.ref = (
        f"A11:{get_column_letter(max_column)}{max(11, total_start - 2)}"
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _composition(payload: dict[str, Any]) -> dict[str, Any]:
    value = payload.get("composition")
    return dict(value) if isinstance(value, dict) else {}


def _enabled_section_types(payload: dict[str, Any]) -> list[str]:
    sections = _composition(payload).get("sections") or []
    enabled = [
        str(item.get("section_type"))
        for item in sorted(
            (item for item in sections if isinstance(item, dict)),
            key=lambda item: int(item.get("position") or 0),
        )
        if item.get("enabled")
    ]
    return enabled or ["SUMMARY", "DETAILED_BOQ", "PAYMENT_TERMS", "TERMS"]


def _write_composed_workbook(
    payload: dict[str, Any],
    *,
    media_bytes: dict[str, bytes],
) -> bytes:
    base = _write_workbook(payload, internal=False)
    workbook = load_workbook(BytesIO(base))
    detailed = workbook.active
    detailed.title = "Detailed BOQ"
    enabled = _enabled_section_types(payload)
    composition = _composition(payload)
    section_by_type = {
        str(item.get("section_type")): item
        for item in composition.get("sections") or []
        if isinstance(item, dict)
    }

    def section_title(section_type: str, fallback: str) -> str:
        item = section_by_type.get(section_type) or {}
        values = (
            str(item.get("title_th") or "").strip(),
            str(item.get("title_en") or "").strip(),
        )
        return " / ".join(value for value in values if value) or fallback

    document = dict(payload.get("document") or {})
    revision = dict(payload.get("revision") or {})
    project = dict(payload.get("project") or {})
    quotation = dict(payload.get("quotation") or {})
    calculation = dict(payload.get("calculation") or {})

    summary = workbook.create_sheet("Summary", 0)
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 52
    rows = [
        (
            section_title("SUMMARY", "สรุปใบเสนอราคา / Quotation Summary"),
            quotation.get("title"),
        ),
        ("เลขที่ / No.", document.get("number")),
        ("ฉบับแก้ไข / Revision", revision.get("number")),
        ("โครงการ / Project", project.get("name")),
        ("ลูกค้า / Customer", quotation.get("customer_name")),
        ("วันที่ / Date", quotation.get("quotation_date")),
        ("ใช้ได้ถึง / Valid until", quotation.get("valid_until")),
        ("รวมก่อน VAT / Net ex VAT", _number(calculation.get("net_sell_ex_vat"))),
        ("VAT", _number(calculation.get("vat_amount"))),
        ("ยอดรวมสุทธิ / Grand total", _number(calculation.get("grand_total"))),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        _set_text(summary.cell(row_index, 1), label)
        summary.cell(row_index, 1).font = Font(bold=True, color="15324A")
        if isinstance(value, (int, float)):
            summary.cell(row_index, 2, value)
            summary.cell(row_index, 2).number_format = '#,##0.00 "THB"'
        else:
            _set_text(summary.cell(row_index, 2), value)
        summary.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    detailed.merge_cells("A10:I10")
    _set_text(
        detailed["A10"],
        section_title("DETAILED_BOQ", "รายละเอียด BOQ / Detailed BOQ"),
    )
    detailed["A10"].font = Font(bold=True, color="15324A")

    if "VISUAL" in enabled:
        visual = workbook.create_sheet("Visuals")
        visual.sheet_view.showGridLines = False
        visual.column_dimensions["A"].width = 4
        visual.column_dimensions["B"].width = 48
        visual.column_dimensions["C"].width = 48
        row_index = 1
        visual.cell(
            row_index,
            1,
            section_title("VISUAL", "รูปภาพและรายละเอียดงาน / Visual & Work Detail"),
        )
        visual.cell(row_index, 1).font = Font(size=16, bold=True, color="15324A")
        visual.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
        row_index += 2
        for page in composition.get("visual_pages") or []:
            if not isinstance(page, dict):
                continue
            visual.cell(row_index, 1, page.get("title_th") or "รูปประกอบ")
            visual.cell(row_index, 1).font = Font(bold=True, color="15324A")
            visual.merge_cells(
                start_row=row_index, start_column=1, end_row=row_index, end_column=3
            )
            row_index += 1
            for entry_index, entry in enumerate(page.get("entries") or []):
                if not isinstance(entry, dict):
                    continue
                media_id = str(entry.get("media_id"))
                image_bytes = media_bytes.get(media_id)
                if image_bytes:
                    image = WorkbookImage(BytesIO(image_bytes))
                    image.width = min(image.width, 360)
                    image.height = min(image.height, 260)
                    column = 2 + (entry_index % 2)
                    visual.add_image(image, f"{get_column_letter(column)}{row_index}")
                caption = " / ".join(
                    value
                    for value in (
                        str(entry.get("caption_th") or "").strip(),
                        str(entry.get("caption_en") or "").strip(),
                    )
                    if value
                )
                visual.cell(row_index + 14, 2 + (entry_index % 2), caption)
                visual.cell(
                    row_index + 14, 2 + (entry_index % 2)
                ).alignment = Alignment(wrap_text=True)
                if entry_index % 2 == 1:
                    row_index += 16
            if len(page.get("entries") or []) % 2:
                row_index += 16
            row_index += 2

    if "PAYMENT_TERMS" in enabled:
        payment = workbook.create_sheet("Payment Terms")
        payment.append(
            [
                section_title("PAYMENT_TERMS", "เงื่อนไขการชำระเงิน / Payment Terms"),
                "จำนวนเงิน / Amount",
            ]
        )
        for cell in payment[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
        for item in quotation.get("payment_schedule") or []:
            if isinstance(item, dict):
                payment.append([item.get("label"), _number(item.get("amount"))])
                payment.cell(payment.max_row, 2).number_format = '#,##0.00 "THB"'
        payment.column_dimensions["A"].width = 60
        payment.column_dimensions["B"].width = 24

    if "TERMS" in enabled or "ACCEPTANCE" in enabled:
        terms = workbook.create_sheet("Terms & Acceptance")
        row_index = 1
        if "TERMS" in enabled:
            terms.cell(
                row_index,
                1,
                section_title("TERMS", "ข้อกำหนดและเงื่อนไข / Terms & Conditions"),
            )
            terms.cell(row_index, 1).font = Font(bold=True, color="15324A")
            row_index += 1
            for index, term in enumerate(
                quotation.get("commercial_terms") or [], start=1
            ):
                terms.cell(row_index, 1, f"{index}. {term}")
                row_index += 1
        if "ACCEPTANCE" in enabled:
            row_index += 2
            terms.cell(
                row_index,
                1,
                section_title("ACCEPTANCE", "การยอมรับใบเสนอราคา / Acceptance"),
            )
            terms.cell(row_index, 1).font = Font(bold=True, color="15324A")
            terms.cell(row_index + 2, 1, "ผู้อนุมัติ / Authorized by: ____________________")
            terms.cell(row_index + 4, 1, "วันที่ / Date: ____________________")
            terms.cell(
                row_index + 6,
                1,
                "Internal recorded agreement — not a digital signature",
            )
        terms.column_dimensions["A"].width = 100

    workbook.active = workbook.sheetnames.index("Detailed BOQ")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_vendor_workbook(payload: dict[str, Any], *, include_prices: bool) -> bytes:
    """Render a server-side allowlisted vendor workbook.

    The payload is frozen at request time and contains neither customer sell
    values nor competing vendor offer data.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Selected Vendor Cost" if include_prices else "RFQ"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A8"
    sheet.print_title_rows = "7:7"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.merge_cells("A1:H1")
    _set_text(sheet["A1"], "Selected Vendor Cost" if include_prices else "Request for Quotation")
    sheet["A1"].font = Font(size=17, bold=True, color="15324A")
    sheet["A1"].alignment = Alignment(horizontal="center")
    identity = [
        ("Project", payload.get("project_name")),
        ("Document", payload.get("document_number")),
        ("Vendor", payload.get("vendor_name")),
        ("Reference", payload.get("quotation_reference") or ""),
    ]
    for index, (label, value) in enumerate(identity, start=2):
        _set_text(sheet.cell(index, 1), label)
        _set_text(sheet.cell(index, 2), value)
        sheet.cell(index, 1).font = Font(bold=True)
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=8)
    columns = ["Item code", "Description", "Specification", "Component", "Quantity", "Unit"]
    if include_prices:
        columns.extend(["Agreed rate ex VAT", "Agreed total ex VAT"])
    else:
        columns.extend(["Vendor rate", "Vendor total"])
    for index, label in enumerate(columns, start=1):
        cell = sheet.cell(7, index)
        _set_text(cell, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row_index, line in enumerate(payload.get("lines") or [], start=8):
        _set_text(sheet.cell(row_index, 1), line.get("item_code"))
        _set_text(sheet.cell(row_index, 2), line.get("description"))
        _set_text(sheet.cell(row_index, 3), line.get("specification"))
        _set_text(sheet.cell(row_index, 4), line.get("component_type"))
        sheet.cell(row_index, 5, _number(line.get("quantity")))
        _set_text(sheet.cell(row_index, 6), line.get("unit"))
        if include_prices:
            sheet.cell(row_index, 7, _number(line.get("unit_rate")))
            sheet.cell(row_index, 8, _number(line.get("total")))
        else:
            sheet.cell(row_index, 7, None)
            sheet.cell(row_index, 8, None)
        for column in range(1, 9):
            sheet.cell(row_index, column).border = Border(bottom=THIN_GREY)
            sheet.cell(row_index, column).alignment = Alignment(
                vertical="top", wrap_text=column in {1, 2, 3, 4, 6},
                horizontal="right" if column in {5, 7, 8} else "left",
            )
        for column in {5, 7, 8}:
            sheet.cell(row_index, column).number_format = "#,##0.00"
    widths = [16, 34, 34, 15, 14, 14, 20, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _snapshot_component_scope(
    snapshot: BOQV2RevisionSnapshot,
) -> dict[UUID, dict[str, Any]]:
    """Return the only fields permitted in a vendor-facing export."""

    internal = snapshot.internal_payload
    if not isinstance(internal, dict):
        raise BOQDomainError(
            "EXPORT_SNAPSHOT_INVALID", "Quotation snapshot payload is invalid"
        )
    result: dict[UUID, dict[str, Any]] = {}
    for node in internal.get("scope") or []:
        if (
            not isinstance(node, dict)
            or node.get("node_kind") != "ITEM"
            or node.get("inclusion_state") == "EXCLUDED"
        ):
            continue
        try:
            scope_node_id = UUID(str(node["id"]))
        except (KeyError, TypeError, ValueError) as error:
            raise BOQDomainError(
                "EXPORT_SNAPSHOT_INVALID", "Snapshot contains an invalid scope identity"
            ) from error
        for component in node.get("components") or []:
            if not isinstance(component, dict):
                continue
            try:
                component_id = UUID(str(component["id"]))
            except (KeyError, TypeError, ValueError) as error:
                raise BOQDomainError(
                    "EXPORT_SNAPSHOT_INVALID",
                    "Snapshot contains an invalid component identity",
                ) from error
            result[component_id] = {
                "component_id": str(component_id),
                "scope_node_id": str(scope_node_id),
                "component_type": str(component.get("component_type") or ""),
                "item_code": _safe_text(node.get("item_code")),
                "description": _safe_text(node.get("description")),
                "specification": _safe_text(
                    component.get("specification") or node.get("specification")
                ),
                "quantity": str(component.get("quantity") or node.get("quantity") or "0"),
                "unit": _safe_text(component.get("unit") or node.get("unit")),
            }
    return result


def _filename_token(value: object | None) -> str:
    token = re.sub(r"[^A-Za-z0-9_-]+", "-", str(value or "vendor")).strip("-")
    return (token or "vendor")[:64]


async def _vendor_export_scope(
    db: AsyncSession,
    *,
    snapshot: BOQV2RevisionSnapshot,
    request: BOQV2ExportRequest,
    project_id: UUID,
    document_number: str,
) -> tuple[BOQV2Vendor, dict[str, Any], int | None]:
    vendor = (
        await db.execute(
            select(BOQV2Vendor).where(
                BOQV2Vendor.id == request.vendor_id,
                BOQV2Vendor.project_id == project_id,
            )
        )
    ).scalar_one_or_none()
    if vendor is None:
        raise BOQDomainError(
            "EXPORT_VENDOR_NOT_FOUND", "Vendor does not belong to this project"
        )

    available = _snapshot_component_scope(snapshot)
    selected_ids = list(request.selected_component_ids)
    if any(component_id not in available for component_id in selected_ids):
        raise BOQDomainError(
            "EXPORT_SCOPE_INVALID",
            "Every selected component must belong to the immutable export snapshot",
        )

    project_payload = snapshot.internal_payload.get("project", {})
    frozen: dict[str, Any] = {
        "schema_version": "boq-v2-vendor-export-v1",
        "project_id": str(project_id),
        "project_name": _safe_text(
            project_payload.get("name") if isinstance(project_payload, dict) else None
        ),
        "document_number": _safe_text(document_number),
        "vendor_id": str(vendor.id),
        "vendor_name": _safe_text(vendor.display_name),
        "quotation_reference": "",
        "lines": [],
    }
    if request.audience == "RFQ":
        frozen["lines"] = [available[component_id] for component_id in selected_ids]
        return vendor, frozen, None

    selection_rows = list(
        (
            await db.execute(
                select(
                    BOQV2CostSelection,
                    BOQV2VendorOfferLine,
                    BOQV2VendorOffer,
                    BOQV2CostPlan,
                )
                .join(
                    BOQV2VendorOfferLine,
                    BOQV2VendorOfferLine.id == BOQV2CostSelection.offer_line_id,
                )
                .join(
                    BOQV2VendorOffer,
                    BOQV2VendorOffer.id == BOQV2VendorOfferLine.offer_id,
                )
                .join(BOQV2CostPlan, BOQV2CostPlan.id == BOQV2CostSelection.cost_plan_id)
                .where(
                    BOQV2CostSelection.project_id == project_id,
                    BOQV2CostSelection.is_current.is_(True),
                    BOQV2CostSelection.status == "CONFIRMED",
                    BOQV2VendorOffer.vendor_id == vendor.id,
                    BOQV2CostPlan.status == "PUBLISHED",
                )
            )
        ).all()
    )
    by_scope = {
        (str(selection.scope_node_id), selection.component_type): (line, offer, plan)
        for selection, line, offer, plan in selection_rows
    }
    if not selected_ids:
        selected_ids = [
            component_id
            for component_id, scope in available.items()
            if (scope["scope_node_id"], scope["component_type"]) in by_scope
        ]
    if not selected_ids:
        raise BOQDomainError(
            "EXPORT_SCOPE_INVALID", "Vendor has no confirmed selection in this snapshot"
        )

    plan_versions: set[int] = set()
    references: set[str] = set()
    lines: list[dict[str, Any]] = []
    for component_id in selected_ids:
        scope = available[component_id]
        selected = by_scope.get((scope["scope_node_id"], scope["component_type"]))
        if selected is None:
            raise BOQDomainError(
                "EXPORT_SCOPE_INVALID",
                "Selected component is not a confirmed cost selection for this vendor",
            )
        line, offer, plan = selected
        if line.normalized_unit_rate_ex_vat is None or line.normalized_total_ex_vat is None:
            raise BOQDomainError(
                "EXPORT_COST_PLAN_MISMATCH",
                "Selected vendor cost has no safe ex-VAT normalized value",
            )
        plan_versions.add(plan.version)
        if offer.quotation_reference:
            references.add(offer.quotation_reference)
        lines.append(
            {
                **scope,
                "unit_rate": str(line.normalized_unit_rate_ex_vat),
                "total": str(line.normalized_total_ex_vat),
                "quotation_reference": _safe_text(offer.quotation_reference),
            }
        )
    if len(plan_versions) != 1:
        raise BOQDomainError(
            "EXPORT_COST_PLAN_MISMATCH",
            "Selected components must come from one published cost-plan version",
        )
    cost_plan_version = plan_versions.pop()
    if (
        request.cost_plan_version is not None
        and request.cost_plan_version != cost_plan_version
    ):
        raise BOQDomainError(
            "EXPORT_COST_PLAN_MISMATCH",
            "Requested cost-plan version does not match the selected vendor costs",
        )
    frozen["quotation_reference"] = ", ".join(sorted(references))
    frozen["cost_plan_version"] = cost_plan_version
    frozen["lines"] = lines
    return vendor, frozen, cost_plan_version


def _font_candidates() -> list[Path]:
    configured = get_settings().boq_pdf_font_path
    paths: list[Path] = []
    if configured:
        paths.append(Path(configured))
    paths.extend(
        [
            Path("/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf"),
            Path("/usr/share/fonts/opentype/noto/NotoSansThai-Regular.ttf"),
            Path.home() / "Library/Fonts/NotoSansThai-Regular.ttf",
        ]
    )
    return paths


def _register_pdf_font() -> str:
    font_name = "BOQNotoSansThai"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    font_path = next((path for path in _font_candidates() if path.is_file()), None)
    if font_path is None:
        raise BOQDomainError(
            "PDF_THAI_FONT_UNAVAILABLE",
            "Noto Sans Thai is required for quotation PDF rendering",
        )
    pdfmetrics.registerFont(TTFont(font_name, str(font_path), shapable=True))
    return font_name


def _write_vendor_pdf(payload: dict[str, Any], *, include_prices: bool) -> bytes:
    font_name = _register_pdf_font()
    output = BytesIO()
    title = "Selected Vendor Cost" if include_prices else "Request for Quotation"
    document_number = _safe_text(payload.get("document_number"))
    vendor_name = _safe_text(payload.get("vendor_name"))
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=17 * mm,
        title=f"{title} {document_number}",
        author="Projects-001",
    )
    base = getSampleStyleSheet()["Normal"]
    normal = ParagraphStyle(
        "VendorNormal",
        parent=base,
        fontName=font_name,
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#243746"),
        wordWrap="CJK",
        shaping=True,
    )
    small = ParagraphStyle("VendorSmall", parent=normal, fontSize=7, leading=9)
    heading = ParagraphStyle(
        "VendorHeading",
        parent=normal,
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#15324A"),
    )

    def paragraph(value: object | None, style: ParagraphStyle = normal) -> Paragraph:
        text = (
            _safe_text(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        return Paragraph(text or " ", style)

    identity = [
        [paragraph("Project"), paragraph(payload.get("project_name"))],
        [paragraph("Document"), paragraph(document_number)],
        [paragraph("Vendor"), paragraph(vendor_name)],
        [paragraph("Reference"), paragraph(payload.get("quotation_reference"))],
    ]
    identity_table = Table(identity, colWidths=[31 * mm, 151 * mm])
    identity_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E7EEF4")),
                ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#C5D0D8")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E0E7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    headers = ["Code", "Description / specification", "Component", "Quantity", "Unit"]
    if include_prices:
        headers.extend(["Agreed rate ex VAT", "Agreed total ex VAT"])
    else:
        headers.extend(["Vendor rate", "Vendor total"])
    table_rows: list[list[Any]] = [[paragraph(value, small) for value in headers]]
    for line in payload.get("lines") or []:
        description = _safe_text(line.get("description"))
        specification = _safe_text(line.get("specification"))
        if specification:
            description = f"{description} — {specification}"
        row = [
            paragraph(line.get("item_code"), small),
            paragraph(description, small),
            paragraph(line.get("component_type"), small),
            paragraph(f"{_decimal(line.get('quantity')):,.4f}", small),
            paragraph(line.get("unit"), small),
        ]
        if include_prices:
            row.extend(
                [
                    paragraph(f"{_decimal(line.get('unit_rate')):,.4f}", small),
                    paragraph(f"{_decimal(line.get('total')):,.2f}", small),
                ]
            )
        else:
            row.extend([paragraph("", small), paragraph("", small)])
        table_rows.append(row)
    details = Table(
        table_rows,
        colWidths=[19 * mm, 65 * mm, 21 * mm, 18 * mm, 16 * mm, 22 * mm, 23 * mm],
        repeatRows=1,
        splitByRow=1,
    )
    details.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#15324A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#AEBBC5")),
                ("INNERGRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D9E0E7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (3, 1), (3, -1), "RIGHT"),
                ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story: list[Any] = [
        paragraph(title, heading),
        Spacer(1, 5 * mm),
        identity_table,
        Spacer(1, 6 * mm),
        details,
    ]

    def footer(canvas: Any, _document: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7)
        canvas.setFillColor(colors.HexColor("#5F6F7A"))
        canvas.drawString(12 * mm, 8 * mm, f"{document_number} · {vendor_name}")
        canvas.drawRightString(A4[0] - 12 * mm, 8 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _write_legacy_pdf(payload: dict[str, Any]) -> bytes:
    font_name = _register_pdf_font()
    output = BytesIO()
    document = dict(payload.get("document") or {})
    revision = dict(payload.get("revision") or {})
    project = dict(payload.get("project") or {})
    quotation = dict(payload.get("quotation") or {})
    calculation = dict(payload.get("calculation") or {})
    doc_number = _safe_text(document.get("number"))
    revision_number = revision.get("number") or 1
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=17 * mm,
        title=f"Quotation {doc_number} R{revision_number}",
        author="Projects-001",
    )
    base = getSampleStyleSheet()["Normal"]
    normal = ParagraphStyle(
        "BOQNormal",
        parent=base,
        fontName=font_name,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#243746"),
        wordWrap="CJK",
        shaping=True,
    )
    small = ParagraphStyle("BOQSmall", parent=normal, fontSize=7.2, leading=9.5)
    heading = ParagraphStyle(
        "BOQHeading",
        parent=normal,
        fontSize=17,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#15324A"),
    )
    right = ParagraphStyle("BOQRight", parent=normal, alignment=TA_RIGHT)

    def paragraph(value: object | None, style: ParagraphStyle = normal) -> Paragraph:
        text = (
            _safe_text(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        return Paragraph(text or " ", style)

    story: list[Any] = [
        paragraph(quotation.get("title") or "ใบเสนอราคา / Quotation", heading),
        Spacer(1, 5 * mm),
    ]
    identity = [
        [paragraph("เลขที่ / No."), paragraph(doc_number), paragraph("Revision"), paragraph(revision_number)],
        [paragraph("โครงการ / Project"), paragraph(project.get("name")), paragraph("วันที่ / Date"), paragraph(quotation.get("quotation_date"))],
        [paragraph("ลูกค้า / Customer"), paragraph(quotation.get("customer_name")), paragraph("ใช้ได้ถึง / Valid until"), paragraph(quotation.get("valid_until"))],
        [paragraph("ที่อยู่ / Address"), paragraph(quotation.get("customer_address")), paragraph("เลขผู้เสียภาษี / Tax ID"), paragraph(quotation.get("customer_tax_id"))],
        [paragraph("ผู้ติดต่อ / Contact"), paragraph(quotation.get("customer_contact")), paragraph("Currency"), paragraph(quotation.get("currency") or "THB")],
    ]
    identity_table = Table(identity, colWidths=[28 * mm, 67 * mm, 31 * mm, 43 * mm])
    identity_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E7EEF4")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E7EEF4")),
                ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#C5D0D8")),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E0E7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([identity_table, Spacer(1, 6 * mm)])

    table_data: list[list[Any]] = [
        [
            paragraph("ลำดับ", small),
            paragraph("รายละเอียด", small),
            paragraph("จำนวน", small),
            paragraph("หน่วย", small),
            paragraph("วัสดุ/หน่วย", small),
            paragraph("แรง/หน่วย", small),
            paragraph("รวม", small),
        ]
    ]
    item_number = 0
    row_kinds: list[str] = []
    for node in _included_scope(payload):
        if node.get("node_kind") != "ITEM":
            table_data.append(
                [paragraph(node.get("description") or node.get("node_kind")), "", "", "", "", "", ""]
            )
            row_kinds.append("section")
            continue
        item_number += 1
        description = node.get("description") or ""
        specification = node.get("specification")
        if specification:
            description = f"{description} — {_safe_text(specification)}"
        table_data.append(
            [
                paragraph(item_number, small),
                paragraph(description, small),
                paragraph(f"{_decimal(node.get('quantity')):,.2f}", small),
                paragraph(node.get("unit"), small),
                paragraph(f"{_decimal(node.get('sell_material_unit_rate')):,.2f}", small),
                paragraph(f"{_decimal(node.get('sell_labor_unit_rate')):,.2f}", small),
                paragraph(f"{_decimal(node.get('sell_total')):,.2f}", small),
            ]
        )
        row_kinds.append("item")
    scope_table = Table(
        table_data,
        colWidths=[12 * mm, 67 * mm, 18 * mm, 15 * mm, 20 * mm, 20 * mm, 22 * mm],
        repeatRows=1,
        splitByRow=1,
    )
    scope_commands: list[tuple[Any, ...]] = [
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#15324A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#AEBBC5")),
        ("INNERGRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D9E0E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for offset, kind in enumerate(row_kinds, start=1):
        if kind == "section":
            scope_commands.extend(
                [
                    ("SPAN", (0, offset), (-1, offset)),
                    ("BACKGROUND", (0, offset), (-1, offset), colors.HexColor("#E7EEF4")),
                ]
            )
    scope_table.setStyle(TableStyle(scope_commands))
    story.extend([scope_table, Spacer(1, 5 * mm)])

    totals = [
        ("รวมก่อนส่วนลด / Subtotal", calculation.get("subtotal")),
        ("ส่วนลด / Discount", calculation.get("discount_amount")),
        ("รวมก่อน VAT / Net", calculation.get("net_sell_ex_vat")),
        (f"VAT {calculation.get('vat_rate') or '0'}%", calculation.get("vat_amount")),
        ("ยอดรวมสุทธิ / Grand total", calculation.get("grand_total")),
    ]
    totals_table = Table(
        [[paragraph(label, right), paragraph(f"{_decimal(value):,.2f} THB", right)] for label, value in totals],
        colWidths=[120 * mm, 49 * mm],
        hAlign="RIGHT",
    )
    totals_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("LINEABOVE", (0, -1), (-1, -1), 0.7, colors.HexColor("#15324A")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EAF2EE")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.extend([totals_table, Spacer(1, 5 * mm)])
    story.append(paragraph("เงื่อนไขการชำระเงิน / Payment schedule"))
    for schedule in quotation.get("payment_schedule") or []:
        if isinstance(schedule, dict):
            story.append(
                paragraph(
                    f"• {schedule.get('label') or ''}: {_decimal(schedule.get('amount')):,.2f} THB",
                    small,
                )
            )
    story.extend([Spacer(1, 3 * mm), paragraph("เงื่อนไขการค้า / Commercial terms")])
    for term in quotation.get("commercial_terms") or []:
        story.append(paragraph(f"• {term}", small))

    def footer(canvas: Any, _document: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7)
        canvas.setFillColor(colors.HexColor("#5F6F7A"))
        canvas.drawString(14 * mm, 8 * mm, f"{doc_number} · R{revision_number}")
        canvas.drawRightString(
            A4[0] - 14 * mm, 8 * mm, f"Page {canvas.getPageNumber()}"
        )
        canvas.restoreState()

    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _write_composed_pdf(
    payload: dict[str, Any],
    *,
    media_bytes: dict[str, bytes],
) -> bytes:
    font_name = _register_pdf_font()
    output = BytesIO()
    document = dict(payload.get("document") or {})
    revision = dict(payload.get("revision") or {})
    project = dict(payload.get("project") or {})
    quotation = dict(payload.get("quotation") or {})
    calculation = dict(payload.get("calculation") or {})
    composition = _composition(payload)
    enabled = _enabled_section_types(payload)
    section_by_type = {
        str(item.get("section_type")): item
        for item in composition.get("sections") or []
        if isinstance(item, dict)
    }
    media_manifest = {
        str(item.get("id")): item
        for item in composition.get("media_assets") or []
        if isinstance(item, dict)
    }
    doc_number = _safe_text(document.get("number"))
    revision_number = revision.get("number") or 1
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=17 * mm,
        title=f"Quotation {doc_number} R{revision_number}",
        author="RAYADEE",
    )
    base = getSampleStyleSheet()["Normal"]
    normal = ParagraphStyle(
        "ComposerNormal",
        parent=base,
        fontName=font_name,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#2F2E2C"),
        wordWrap="CJK",
        shaping=True,
    )
    small = ParagraphStyle("ComposerSmall", parent=normal, fontSize=7.5, leading=10)
    title_style = ParagraphStyle(
        "ComposerTitle",
        parent=normal,
        fontSize=21,
        leading=27,
        textColor=colors.HexColor("#2F2E2C"),
    )
    heading = ParagraphStyle(
        "ComposerHeading",
        parent=normal,
        fontSize=14,
        leading=19,
        textColor=colors.HexColor("#4F6F64"),
    )
    right = ParagraphStyle("ComposerRight", parent=normal, alignment=TA_RIGHT)

    def paragraph(value: object | None, style: ParagraphStyle = normal) -> Paragraph:
        text = (
            _safe_text(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        return Paragraph(text or " ", style)

    def section_title(section_type: str) -> str:
        item = section_by_type.get(section_type) or {}
        th = str(item.get("title_th") or "").strip()
        en = str(item.get("title_en") or "").strip()
        return (
            " / ".join(value for value in (th, en) if value)
            or section_type.replace("_", " ").title()
        )

    def totals_table() -> Table:
        rows = [
            ("รวมก่อนส่วนลด / Subtotal", calculation.get("subtotal")),
            ("ส่วนลด / Discount", calculation.get("discount_amount")),
            ("รวมก่อน VAT / Net ex VAT", calculation.get("net_sell_ex_vat")),
            (
                f"VAT {calculation.get('vat_rate') or '0'}%",
                calculation.get("vat_amount"),
            ),
            ("ยอดรวมสุทธิ / Grand total", calculation.get("grand_total")),
        ]
        table = Table(
            [
                [
                    paragraph(label, right),
                    paragraph(f"{_decimal(value):,.2f} THB", right),
                ]
                for label, value in rows
            ],
            colWidths=[112 * mm, 57 * mm],
            hAlign="RIGHT",
        )
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("LINEABOVE", (0, -1), (-1, -1), 0.7, colors.HexColor("#4F6F64")),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EAF2EE")),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return table

    story: list[Any] = []
    started = False

    def begin_section(section_type: str) -> None:
        nonlocal started
        if started:
            story.append(PageBreak())
        started = True
        story.extend(
            [paragraph(section_title(section_type), heading), Spacer(1, 5 * mm)]
        )

    for section_type in enabled:
        if section_type == "SUMMARY":
            begin_section(section_type)
            story.extend(
                [
                    paragraph(
                        quotation.get("title") or "ใบเสนอราคา / Quotation", title_style
                    ),
                    Spacer(1, 6 * mm),
                ]
            )
            identity = [
                [
                    paragraph("เลขที่ / No."),
                    paragraph(doc_number),
                    paragraph("ฉบับ / Revision"),
                    paragraph(revision_number),
                ],
                [
                    paragraph("โครงการ / Project"),
                    paragraph(project.get("name")),
                    paragraph("วันที่ / Date"),
                    paragraph(quotation.get("quotation_date")),
                ],
                [
                    paragraph("ลูกค้า / Customer"),
                    paragraph(quotation.get("customer_name")),
                    paragraph("ใช้ได้ถึง / Valid until"),
                    paragraph(quotation.get("valid_until")),
                ],
                [
                    paragraph("ที่อยู่ / Address"),
                    paragraph(quotation.get("customer_address")),
                    paragraph("เลขผู้เสียภาษี / Tax ID"),
                    paragraph(quotation.get("customer_tax_id")),
                ],
                [
                    paragraph("ผู้ติดต่อ / Contact"),
                    paragraph(quotation.get("customer_contact")),
                    paragraph("สกุลเงิน / Currency"),
                    paragraph(quotation.get("currency") or "THB"),
                ],
            ]
            table = Table(identity, colWidths=[29 * mm, 65 * mm, 32 * mm, 43 * mm])
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), font_name),
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F0EC")),
                        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F3F0EC")),
                        (
                            "LINEBELOW",
                            (0, 0),
                            (-1, -1),
                            0.25,
                            colors.HexColor("#D8D2CA"),
                        ),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.extend([table, Spacer(1, 10 * mm), totals_table()])

        elif section_type == "DETAILED_BOQ":
            begin_section(section_type)
            table_data: list[list[Any]] = [
                [
                    paragraph("#", small),
                    paragraph("รายละเอียด / Scope", small),
                    paragraph("จำนวน / Qty", small),
                    paragraph("หน่วย / Unit", small),
                    paragraph("วัสดุ / Material", small),
                    paragraph("แรง / Labor", small),
                    paragraph("รวม / Total", small),
                ]
            ]
            row_kinds: list[str] = []
            item_number = 0
            for node in _included_scope(payload):
                if node.get("node_kind") != "ITEM":
                    table_data.append(
                        [
                            paragraph(node.get("description") or node.get("node_kind")),
                            "",
                            "",
                            "",
                            "",
                            "",
                            "",
                        ]
                    )
                    row_kinds.append("section")
                    continue
                item_number += 1
                description = _safe_text(node.get("description"))
                if node.get("specification"):
                    description = f"{description}<br/><font size='7'>{_safe_text(node.get('specification'))}</font>"
                table_data.append(
                    [
                        paragraph(item_number, small),
                        paragraph(description, small),
                        paragraph(f"{_decimal(node.get('quantity')):,.2f}", small),
                        paragraph(node.get("unit"), small),
                        paragraph(
                            f"{_decimal(node.get('sell_material_unit_rate')):,.2f}",
                            small,
                        ),
                        paragraph(
                            f"{_decimal(node.get('sell_labor_unit_rate')):,.2f}", small
                        ),
                        paragraph(f"{_decimal(node.get('sell_total')):,.2f}", small),
                    ]
                )
                row_kinds.append("item")
            scope_table = Table(
                table_data,
                colWidths=[
                    10 * mm,
                    70 * mm,
                    17 * mm,
                    16 * mm,
                    20 * mm,
                    18 * mm,
                    22 * mm,
                ],
                repeatRows=1,
                splitByRow=1,
            )
            commands: list[tuple[Any, ...]] = [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F6F64")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8D2CA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (2, 1), (2, -1), "RIGHT"),
                ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
            for offset, kind in enumerate(row_kinds, start=1):
                if kind == "section":
                    commands.extend(
                        [
                            ("SPAN", (0, offset), (-1, offset)),
                            (
                                "BACKGROUND",
                                (0, offset),
                                (-1, offset),
                                colors.HexColor("#EDF2F0"),
                            ),
                        ]
                    )
            scope_table.setStyle(TableStyle(commands))
            story.extend([scope_table, Spacer(1, 6 * mm), totals_table()])

        elif section_type == "VISUAL":
            pages = [
                page
                for page in composition.get("visual_pages") or []
                if isinstance(page, dict)
            ]
            if not pages:
                continue
            begin_section(section_type)
            for page_index, page in enumerate(pages):
                if page_index:
                    story.append(PageBreak())
                page_title = " / ".join(
                    value
                    for value in (
                        str(page.get("title_th") or "").strip(),
                        str(page.get("title_en") or "").strip(),
                    )
                    if value
                )
                if page_title:
                    story.extend([paragraph(page_title, heading), Spacer(1, 2 * mm)])
                descriptions = " / ".join(
                    value
                    for value in (
                        str(page.get("description_th") or "").strip(),
                        str(page.get("description_en") or "").strip(),
                    )
                    if value
                )
                if descriptions:
                    story.extend([paragraph(descriptions), Spacer(1, 4 * mm)])
                layout = str(page.get("layout") or "TWO_UP")
                columns = 1 if layout == "SINGLE" else 2
                cell_width = 169 * mm if columns == 1 else 82 * mm
                max_height = (
                    112 * mm
                    if columns == 1
                    else (72 * mm if layout == "TWO_UP" else 55 * mm)
                )
                cells: list[Any] = []
                for entry in page.get("entries") or []:
                    if not isinstance(entry, dict):
                        continue
                    media_id = str(entry.get("media_id"))
                    image_bytes = media_bytes.get(media_id)
                    manifest = media_manifest.get(media_id) or {}
                    if not image_bytes:
                        raise BOQDomainError(
                            "EXPORT_MEDIA_MISSING",
                            "Frozen quotation visual media is unavailable",
                        )
                    visual = ReportLabImage(BytesIO(image_bytes))
                    width = float(manifest.get("width") or visual.imageWidth or 1)
                    height = float(manifest.get("height") or visual.imageHeight or 1)
                    scale = min(cell_width / width, max_height / height)
                    visual.drawWidth = width * scale
                    visual.drawHeight = height * scale
                    caption = " / ".join(
                        value
                        for value in (
                            str(entry.get("caption_th") or "").strip(),
                            str(entry.get("caption_en") or "").strip(),
                        )
                        if value
                    )
                    cell = Table(
                        [[visual], [paragraph(caption or " ", small)]],
                        colWidths=[cell_width],
                    )
                    cell.setStyle(
                        TableStyle(
                            [
                                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ]
                        )
                    )
                    cells.append(cell)
                rows = [
                    cells[index : index + columns]
                    for index in range(0, len(cells), columns)
                ]
                if rows:
                    if len(rows[-1]) < columns:
                        rows[-1].extend([""] * (columns - len(rows[-1])))
                    gallery = Table(
                        rows, colWidths=[cell_width] * columns, hAlign="LEFT"
                    )
                    gallery.setStyle(
                        TableStyle(
                            [
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                            ]
                        )
                    )
                    story.append(gallery)

        elif section_type == "PAYMENT_TERMS":
            begin_section(section_type)
            schedule_rows = [
                [paragraph("งวด / Milestone"), paragraph("จำนวนเงิน / Amount", right)]
            ]
            for item in quotation.get("payment_schedule") or []:
                if isinstance(item, dict):
                    schedule_rows.append(
                        [
                            paragraph(item.get("label")),
                            paragraph(
                                f"{_decimal(item.get('amount')):,.2f} THB", right
                            ),
                        ]
                    )
            table = Table(schedule_rows, colWidths=[115 * mm, 54 * mm])
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), font_name),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F6F64")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        (
                            "LINEBELOW",
                            (0, 0),
                            (-1, -1),
                            0.25,
                            colors.HexColor("#D8D2CA"),
                        ),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.append(table)

        elif section_type == "TERMS":
            begin_section(section_type)
            for index, term in enumerate(
                quotation.get("commercial_terms") or [], start=1
            ):
                story.extend([paragraph(f"{index}. {term}"), Spacer(1, 2 * mm)])

        elif section_type == "ACCEPTANCE":
            begin_section(section_type)
            story.extend(
                [
                    paragraph(
                        "ข้าพเจ้ายอมรับรายละเอียด ราคา และเงื่อนไขตามใบเสนอราคาฉบับนี้ / I accept the scope, price, and terms of this quotation."
                    ),
                    Spacer(1, 22 * mm),
                    paragraph(
                        "ลงชื่อผู้มีอำนาจ / Authorized signature: ______________________________"
                    ),
                    Spacer(1, 12 * mm),
                    paragraph("ชื่อ / Name: ______________________________"),
                    Spacer(1, 12 * mm),
                    paragraph("วันที่ / Date: ______________________________"),
                    Spacer(1, 18 * mm),
                    paragraph(
                        "บันทึกการยอมรับภายใน ไม่ใช่ลายมือชื่อดิจิทัล / Internal recorded agreement — not a digital signature",
                        small,
                    ),
                ]
            )

    def footer(canvas: Any, _document: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7)
        canvas.setFillColor(colors.HexColor("#5F6F7A"))
        canvas.drawString(14 * mm, 8 * mm, f"{doc_number} · R{revision_number}")
        canvas.drawRightString(
            A4[0] - 14 * mm, 8 * mm, f"Page {canvas.getPageNumber()}"
        )
        canvas.restoreState()

    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_export_bytes(
    snapshot: BOQV2RevisionSnapshot,
    *,
    audience: str,
    file_format: str,
    audience_payload: dict[str, Any] | None = None,
    media_bytes: dict[str, bytes] | None = None,
) -> bytes:
    if audience in {"RFQ", "VENDOR"}:
        if not isinstance(audience_payload, dict):
            raise BOQDomainError(
                "EXPORT_SCOPE_INVALID", "Frozen vendor export scope is unavailable"
            )
        if file_format == "XLSX":
            return _write_vendor_workbook(
                audience_payload, include_prices=audience == "VENDOR"
            )
        if file_format == "PDF":
            return _write_vendor_pdf(
                audience_payload, include_prices=audience == "VENDOR"
            )
    payload = _payload(snapshot, audience)
    if file_format == "XLSX":
        if (
            audience == "CUSTOMER"
            and payload.get("schema_version") == "boq-v2-document-snapshot-v2"
        ):
            return _write_composed_workbook(payload, media_bytes=media_bytes or {})
        return _write_workbook(payload, internal=audience == "INTERNAL")
    if file_format == "PDF" and audience == "CUSTOMER":
        if payload.get("schema_version") == "boq-v2-document-snapshot-v2":
            return _write_composed_pdf(payload, media_bytes=media_bytes or {})
        return _write_legacy_pdf(payload)
    raise BOQDomainError(
        "EXPORT_FORMAT_NOT_SUPPORTED",
        "Unsupported export audience/format combination",
    )


def _artifact_response(
    artifact: BOQV2ExportArtifact,
    *,
    download_url: str | None = None,
    expires_in_minutes: int | None = None,
) -> BOQV2ExportArtifactResponse:
    return BOQV2ExportArtifactResponse(
        artifact_id=artifact.id,
        project_id=artifact.project_id,
        document_id=artifact.document_id,
        revision_id=artifact.revision_id,
        snapshot_id=artifact.snapshot_id,
        calculation_version=artifact.calculation_version,
        audience=artifact.audience,
        file_format=artifact.file_format,
        cost_plan_version=artifact.cost_plan_version,
        vendor_id=artifact.vendor_id,
        status=artifact.status,
        filename=artifact.filename,
        mime_type=artifact.mime_type,
        sha256=artifact.sha256,
        size_bytes=artifact.size_bytes,
        created_at=_timestamp(artifact.created_at),
        completed_at=_timestamp(artifact.completed_at) if artifact.completed_at else None,
        download_url=download_url,
        download_expires_in_minutes=expires_in_minutes,
    )


async def request_export(
    db: AsyncSession,
    *,
    revision_id: UUID,
    request: BOQV2ExportRequest,
    actor: str,
    idempotency_key: str,
) -> BOQV2ExportArtifact:
    revision, document, _ = await _revision_context(db, revision_id)
    snapshot = (
        await db.execute(
            select(BOQV2RevisionSnapshot).where(
                BOQV2RevisionSnapshot.id == request.snapshot_id,
                BOQV2RevisionSnapshot.revision_id == revision.id,
                BOQV2RevisionSnapshot.document_id == document.id,
                BOQV2RevisionSnapshot.project_id == revision.project_id,
            )
        )
    ).scalar_one_or_none()
    if snapshot is None:
        raise BOQDomainError(
            "QUOTATION_SNAPSHOT_NOT_FOUND",
            "Quotation snapshot does not belong to this revision",
        )
    command, prior_artifact_id = await _begin_command(
        db,
        actor=actor,
        project_id=revision.project_id,
        command_type="EXPORT_QUOTATION",
        idempotency_key=idempotency_key,
        payload={
            "revision_id": str(revision.id),
            "snapshot_id": str(snapshot.id),
            "audience": request.audience,
            "file_format": request.file_format,
            "vendor_id": str(request.vendor_id) if request.vendor_id else None,
            "selected_component_ids": [str(item) for item in request.selected_component_ids],
            "cost_plan_version": request.cost_plan_version,
        },
    )
    if prior_artifact_id is not None:
        artifact = (
            await db.execute(
                select(BOQV2ExportArtifact).where(
                    BOQV2ExportArtifact.id == prior_artifact_id,
                    BOQV2ExportArtifact.revision_id == revision.id,
                )
            )
        ).scalar_one_or_none()
        if artifact is None:
            raise BOQDomainError(
                "EXPORT_ARTIFACT_NOT_FOUND",
                "Idempotent export artifact is unavailable",
            )
        return artifact

    vendor: BOQV2Vendor | None = None
    audience_payload: dict[str, Any] | None = None
    vendor_cost_plan_version: int | None = None
    if request.audience in {"RFQ", "VENDOR"}:
        vendor, audience_payload, vendor_cost_plan_version = await _vendor_export_scope(
            db,
            snapshot=snapshot,
            request=request,
            project_id=revision.project_id,
            document_number=document.document_number,
        )

    frozen_cost = (
        snapshot.internal_payload.get("cost", {})
        if isinstance(snapshot.internal_payload, dict)
        else {}
    )
    frozen_cost_plan_version = frozen_cost.get("plan_version")
    frozen_cost_plan_id = frozen_cost.get("plan_id")
    if request.audience == "INTERNAL" and (
        (frozen_cost_plan_id is None) != (frozen_cost_plan_version is None)
    ):
        raise BOQDomainError(
            "EXPORT_COST_PLAN_MISMATCH",
            "Frozen internal snapshot has an incomplete cost plan identity",
        )
    extension = request.file_format.lower()
    if request.audience == "CUSTOMER":
        suffix = "customer"
    elif request.audience == "INTERNAL":
        suffix = "internal"
    else:
        suffix = f"{request.audience.casefold()}-{_filename_token(vendor.display_name if vendor else None)}"
    filename = f"{document.document_number}-R{revision.revision_number}-{suffix}.{extension}"
    artifact = BOQV2ExportArtifact(
        id=uuid4(),
        project_id=revision.project_id,
        document_id=document.id,
        revision_id=revision.id,
        snapshot_id=snapshot.id,
        calculation_version=snapshot.calculation_version,
        audience=request.audience,
        file_format=request.file_format,
        cost_plan_version=(
            int(frozen_cost_plan_version)
            if request.audience == "INTERNAL" and frozen_cost_plan_version is not None
            else vendor_cost_plan_version
        ),
        vendor_id=vendor.id if vendor else None,
        audience_payload=audience_payload,
        status="PENDING",
        filename=filename,
        mime_type=XLSX_MIME if request.file_format == "XLSX" else PDF_MIME,
        created_by=actor,
    )
    db.add(artifact)
    await db.flush()
    _complete_command(command, revision_id=artifact.id, version=snapshot.source_version)
    await db.flush()
    return artifact


async def load_export_artifact(
    db: AsyncSession,
    *,
    revision_id: UUID,
    artifact_id: UUID,
    for_update: bool = False,
) -> BOQV2ExportArtifact:
    await _revision_context(db, revision_id)
    statement = select(BOQV2ExportArtifact).where(
        BOQV2ExportArtifact.id == artifact_id,
        BOQV2ExportArtifact.revision_id == revision_id,
    )
    if for_update:
        statement = statement.with_for_update()
    artifact = (await db.execute(statement)).scalar_one_or_none()
    if artifact is None:
        raise BOQDomainError(
            "EXPORT_ARTIFACT_NOT_FOUND",
            "Export artifact does not belong to this revision",
        )
    return artifact


async def render_and_upload_export(
    db: AsyncSession,
    *,
    revision_id: UUID,
    artifact_id: UUID,
) -> BOQV2ExportArtifactResponse:
    artifact = await load_export_artifact(
        db, revision_id=revision_id, artifact_id=artifact_id, for_update=True
    )
    if artifact.status == "READY" and artifact.storage_key:
        return _artifact_response(artifact)
    snapshot = (
        await db.execute(
            select(BOQV2RevisionSnapshot).where(
                BOQV2RevisionSnapshot.id == artifact.snapshot_id,
                BOQV2RevisionSnapshot.revision_id == revision_id,
            )
        )
    ).scalar_one_or_none()
    if snapshot is None:
        artifact.status = "FAILED"
        artifact.error_code = "EXPORT_SNAPSHOT_MISSING"
        await db.flush()
        raise BOQDomainError(
            "EXPORT_SNAPSHOT_MISSING", "Immutable export snapshot is unavailable"
        )
    try:
        media_bytes: dict[str, bytes] = {}
        if (
            artifact.audience == "CUSTOMER"
            and dict(snapshot.customer_payload or {}).get("schema_version")
            == "boq-v2-document-snapshot-v2"
        ):
            composition = dict(
                dict(snapshot.internal_payload or {}).get("composition") or {}
            )
            for item in composition.get("media_assets") or []:
                if not isinstance(item, dict):
                    continue
                media_id = str(item.get("id") or "")
                storage_key = str(item.get("storage_key") or "")
                expected_hash = str(item.get("sha256") or "")
                if not media_id or not storage_key or not expected_hash:
                    raise BOQDomainError(
                        "EXPORT_MEDIA_MISSING",
                        "Frozen quotation visual media metadata is incomplete",
                    )
                rendered_bytes = await download_storage_key_bytes(storage_key)
                if hashlib.sha256(rendered_bytes).hexdigest() != expected_hash:
                    raise BOQDomainError(
                        "EXPORT_MEDIA_HASH_MISMATCH",
                        "Frozen quotation visual media failed its integrity check",
                    )
                media_bytes[media_id] = rendered_bytes
        file_bytes = render_export_bytes(
            snapshot,
            audience=artifact.audience,
            file_format=artifact.file_format,
            audience_payload=artifact.audience_payload,
            media_bytes=media_bytes,
        )
        storage_key = await upload_boq_export_artifact(
            project_id=str(artifact.project_id),
            document_id=str(artifact.document_id),
            revision_id=str(artifact.revision_id),
            snapshot_id=str(artifact.snapshot_id),
            artifact_id=str(artifact.id),
            file_name=artifact.filename,
            file_bytes=file_bytes,
            content_type=artifact.mime_type,
        )
    except BOQDomainError as error:
        artifact.status = "FAILED"
        artifact.error_code = "EXPORT_RENDER_FAILED"
        artifact.completed_at = _now()
        await db.flush()
        raise BOQDomainError(
            "EXPORT_RENDER_FAILED", "Export could not be rendered; retry safely"
        ) from error
    except Exception as error:
        artifact.status = "FAILED"
        artifact.error_code = "EXPORT_STORAGE_FAILED"
        artifact.completed_at = _now()
        await db.flush()
        raise BOQDomainError(
            "EXPORT_STORAGE_FAILED", "Export could not be stored; retry safely"
        ) from error

    artifact.status = "READY"
    artifact.storage_key = storage_key
    artifact.sha256 = hashlib.sha256(file_bytes).hexdigest()
    artifact.size_bytes = len(file_bytes)
    artifact.error_code = None
    artifact.completed_at = _now()
    await db.flush()
    return _artifact_response(artifact)


async def load_export_response(
    db: AsyncSession,
    *,
    revision_id: UUID,
    artifact_id: UUID,
) -> BOQV2ExportArtifactResponse:
    artifact = await load_export_artifact(
        db, revision_id=revision_id, artifact_id=artifact_id
    )
    return _artifact_response(artifact)


async def create_export_download(
    db: AsyncSession,
    *,
    revision_id: UUID,
    artifact_id: UUID,
    expires_in_minutes: int = 15,
) -> BOQV2ExportArtifactResponse:
    artifact = await load_export_artifact(
        db, revision_id=revision_id, artifact_id=artifact_id
    )
    if artifact.status != "READY" or not artifact.storage_key:
        raise BOQDomainError(
            "EXPORT_NOT_READY", "Export artifact is not ready for download"
        )
    url = await generate_signed_url_for_storage_key(
        storage_key=artifact.storage_key,
        expires_in_minutes=expires_in_minutes,
    )
    return _artifact_response(
        artifact,
        download_url=url,
        expires_in_minutes=expires_in_minutes,
    )
