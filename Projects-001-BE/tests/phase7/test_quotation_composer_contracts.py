from __future__ import annotations

import subprocess
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from PIL import Image
from pydantic import ValidationError

from app.schemas.boq_v2_schema import BOQV2QuotationDraft
from app.core.config import get_settings
from app.services.boq_domain_service import BOQDomainError
from app.services.boq_composition_service import default_sections
from app.services.boq_export_service import render_export_bytes
from app.services.boq_quotation_media_service import _normalize_image


def _image_bytes(color: str = "#4f6f64") -> bytes:
    output = BytesIO()
    Image.new("RGB", (640, 420), color).save(output, format="PNG")
    return output.getvalue()


def _snapshot() -> tuple[SimpleNamespace, dict[str, bytes]]:
    media_id = str(uuid4())
    image_bytes, width, height, digest = _normalize_image(_image_bytes(), "image/png")
    sections = [item.model_dump(mode="json") for item in default_sections()]
    for item in sections:
        if item["section_type"] in {"VISUAL", "ACCEPTANCE"}:
            item["enabled"] = True
    common = {
        "schema_version": "boq-v2-document-snapshot-v2",
        "document": {"number": "QT-2026-000777"},
        "revision": {"number": 2},
        "project": {"name": "โครงการทดสอบ Phase 7"},
        "quotation": {
            "title": "Renovation / งานปรับปรุง",
            "customer_name": "บริษัท ลูกค้าทดสอบ จำกัด",
            "customer_address": "กรุงเทพมหานคร",
            "customer_tax_id": "0105559999999",
            "customer_contact": "คุณทดสอบ",
            "quotation_date": "2026-09-22",
            "valid_until": "2026-10-22",
            "currency": "THB",
            "payment_schedule": [
                {"label": "มัดจำ / Mobilization", "amount": "535.00"},
                {"label": "ส่งมอบ / Handover", "amount": "535.00"},
            ],
            "commercial_terms": ["ยืนราคา 30 วัน / Valid for 30 days"],
        },
        "calculation": {
            "subtotal": "1000.00",
            "discount_amount": "0.00",
            "net_sell_ex_vat": "1000.00",
            "vat_rate": "7.0000",
            "vat_amount": "70.00",
            "grand_total": "1070.00",
        },
        "composition": {
            "schema_version": "boq-v2-document-snapshot-v2",
            "sections": sections,
            "visual_pages": [
                {
                    "id": str(uuid4()),
                    "position": 0,
                    "layout": "SINGLE",
                    "title_th": "รูปหน้างาน",
                    "title_en": "Site condition",
                    "description_th": "ก่อนเริ่มงาน",
                    "description_en": "Before construction",
                    "entries": [
                        {
                            "id": str(uuid4()),
                            "media_id": media_id,
                            "position": 0,
                            "caption_th": "พื้นที่ติดตั้ง",
                            "caption_en": "Installation area",
                        }
                    ],
                }
            ],
            "media_assets": [
                {
                    "id": media_id,
                    "origin_type": "UPLOAD",
                    "original_filename": "site.png",
                    "content_type": "image/jpeg",
                    "size_bytes": len(image_bytes),
                    "width": width,
                    "height": height,
                    "sha256": digest,
                }
            ],
        },
        "scope": [
            {
                "id": str(uuid4()),
                "logical_id": str(uuid4()),
                "parent_logical_id": None,
                "node_kind": "ITEM",
                "inclusion_state": "REQUIRED",
                "position": 0,
                "depth": 0,
                "item_code": "FL-001",
                "description": "งานพื้น / Flooring",
                "specification": "SPC 5 mm",
                "quantity": "10.0000",
                "unit": "m²",
                "sell_material_unit_rate": "80.0000",
                "sell_labor_unit_rate": "20.0000",
                "sell_total": "1000.00",
            }
        ],
    }
    return SimpleNamespace(customer_payload=common, internal_payload=common), {
        media_id: image_bytes
    }


def test_section_contract_enforces_summary_boq_and_acceptance_order() -> None:
    sections = default_sections()
    accepted = [item.model_copy() for item in sections]
    accepted[-1].enabled = True
    draft = BOQV2QuotationDraft(document_sections=accepted)
    assert draft.document_schema_version == "boq-v2-document-snapshot-v2"

    invalid = [item.model_copy() for item in sections]
    invalid[0].enabled = False
    with pytest.raises(ValidationError, match="SUMMARY must be enabled"):
        BOQV2QuotationDraft(document_sections=invalid)


def test_image_normalization_is_jpeg_and_deterministically_hashed() -> None:
    normalized, width, height, digest = _normalize_image(_image_bytes(), "image/png")
    assert normalized.startswith(b"\xff\xd8\xff")
    assert (width, height) == (640, 420)
    assert len(digest) == 64
    with Image.open(BytesIO(normalized)) as image:
        assert not image.getexif()

    with pytest.raises(BOQDomainError, match="corrupt or cannot be decoded"):
        _normalize_image(_image_bytes(), "image/jpeg")

    with pytest.raises(BOQDomainError, match="corrupt or cannot be decoded"):
        _normalize_image(b"not-an-image", "image/png")


def test_image_normalization_rejects_oversized_input(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_quotation_media_max_bytes", 32)
    with pytest.raises(BOQDomainError, match="exceeds"):
        _normalize_image(_image_bytes(), "image/png")


def test_v2_customer_exports_include_sections_visuals_and_nonblank_pdf(
    tmp_path,
) -> None:
    snapshot, media_bytes = _snapshot()
    workbook_bytes = render_export_bytes(
        snapshot,
        audience="CUSTOMER",
        file_format="XLSX",
        media_bytes=media_bytes,
    )
    workbook = load_workbook(BytesIO(workbook_bytes))
    assert workbook.sheetnames == [
        "Summary",
        "Detailed BOQ",
        "Visuals",
        "Payment Terms",
        "Terms & Acceptance",
    ]
    assert len(workbook["Visuals"]._images) == 1
    assert workbook["Summary"]["B10"].value == 1070
    visual_text = " ".join(
        str(cell.value or "") for row in workbook["Visuals"].iter_rows() for cell in row
    )
    assert "Installation area" in visual_text
    assert "storage_key" not in visual_text

    pdf_bytes = render_export_bytes(
        snapshot,
        audience="CUSTOMER",
        file_format="PDF",
        media_bytes=media_bytes,
    )
    pdf_path = tmp_path / "phase7-quotation.pdf"
    pdf_path.write_bytes(pdf_bytes)
    info = subprocess.run(
        ["pdfinfo", str(pdf_path)],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    pages_line = next(line for line in info.splitlines() if line.startswith("Pages:"))
    assert int(pages_line.split(":", 1)[1].strip()) >= 6
    extracted = subprocess.run(
        ["pdftotext", str(pdf_path), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    assert "QT-2026-000777" in extracted
    assert "Detailed BOQ" in extracted
    assert "Installation area" in extracted
    assert "Internal recorded agreement" in extracted
    render_prefix = tmp_path / "phase7-page"
    subprocess.run(
        [
            "pdftoppm",
            "-f",
            "1",
            "-singlefile",
            "-png",
            str(pdf_path),
            str(render_prefix),
        ],
        check=True,
        capture_output=True,
    )
    assert render_prefix.with_suffix(".png").stat().st_size > 10_000
