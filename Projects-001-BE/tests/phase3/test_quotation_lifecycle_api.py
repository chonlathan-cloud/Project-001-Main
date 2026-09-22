from __future__ import annotations

import asyncio
import hashlib
import os
from datetime import date
from io import BytesIO
from uuid import UUID, uuid4

import asyncpg
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.core.database import engine
from app.schemas.boq_quotation_schema import BOQV2RecordAcceptanceRequest
from app.services.boq_domain_service import BOQDomainError
from app.services.boq_numbering_service import next_document_number
from app.services.boq_quotation_service import record_acceptance
from main import app


pytestmark = pytest.mark.skipif(
    not os.getenv("PHASE3_DATABASE_URL"),
    reason="PHASE3_DATABASE_URL isolated PostgreSQL target is required",
)

OWNER_HEADERS = {"X-Debug-Admin-Role": "owner"}
ADMIN_HEADERS = {"X-Debug-Admin-Role": "admin"}


def _key(prefix: str) -> str:
    return f"phase3-{prefix}-{uuid4()}"


def _project(client: TestClient, label: str) -> str:
    response = client.post(
        "/api/v1/projects",
        headers=OWNER_HEADERS,
        json={
            "name": f"Phase 3 {label} {uuid4()}",
            "project_type": "COMMERCIAL",
            "overhead_percent": 0,
            "profit_percent": 0,
            "vat_percent": 7,
            "contingency_budget": 25000,
            "status": "ACTIVE",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]["project_id"]


def _quotation(customer: str = "บริษัท ลูกค้าทดสอบ จำกัด") -> dict[str, object]:
    return {
        "title": "ใบเสนอราคาโครงการทดสอบ",
        "customer_name": customer,
        "customer_address": "99 ถนนสุขุมวิท กรุงเทพมหานคร",
        "customer_tax_id": "0105559999999",
        "customer_contact": "คุณทดสอบ",
        "quotation_date": "2026-09-17",
        "valid_until": "2026-10-17",
        "currency": "THB",
        "vat_rate": "7.0000",
        "discount_type": "NONE",
        "discount_value": "0.0000",
        "payment_schedule": [
            {"label": "มัดจำ", "percentage": "50.0000"},
            {"label": "ส่งมอบ", "percentage": "50.0000"},
        ],
        "commercial_terms": ["ยืนราคา 30 วัน", "ชำระตามงวดที่ระบุ"],
        "document_pages": ["BOQ", "PAYMENT_TERMS", "COMMERCIAL_TERMS"],
    }


def _scope_item(
    *,
    row_id: str | None = None,
    logical_id: str | None = None,
    description: str = "งานติดตั้งระบบไฟฟ้า",
    quantity: str = "10.0000",
    material: str = "800.0000",
    labor: str = "200.0000",
) -> dict[str, object]:
    return {
        "id": row_id,
        "logical_id": logical_id or str(uuid4()),
        "parent_logical_id": None,
        "node_kind": "ITEM",
        "inclusion_state": "REQUIRED",
        "position": 0,
        "item_code": "=FORMULA-GUARD",
        "description": description,
        "specification": "มาตรฐานงานก่อสร้างภาษาไทย",
        "quantity": quantity,
        "unit": "งาน",
        "sell_material_unit_rate": material,
        "sell_labor_unit_rate": labor,
        "components": [
            {
                "component_type": "MATERIAL",
                "quantity_basis": "INHERITED",
                "cost_state": "PRICED",
                "unit_rate": "500.0000",
            },
            {
                "component_type": "LABOR",
                "quantity_basis": "INHERITED",
                "cost_state": "PRICED",
                "unit_rate": "100.0000",
            },
        ],
    }


def _create_saved_main(
    client: TestClient,
    project_id: str,
    *,
    customer: str = "บริษัท ลูกค้าทดสอบ จำกัด",
    amount: str = "800.0000",
    logical_id: str | None = None,
) -> dict[str, object]:
    created = client.post(
        f"/api/v1/projects/{project_id}/boq/documents",
        headers={**OWNER_HEADERS, "Idempotency-Key": _key("create")},
        json={},
    )
    assert created.status_code == 201, created.text
    revision = created.json()["data"]
    saved = client.patch(
        f"/api/v1/boq/revisions/{revision['revision_id']}",
        headers={**OWNER_HEADERS, "Idempotency-Key": _key("save")},
        json={
            "expected_version": revision["version"],
            "quotation": _quotation(customer),
            "nodes": [_scope_item(material=amount, logical_id=logical_id)],
        },
    )
    assert saved.status_code == 200, saved.text
    return saved.json()["data"]


def _issue(client: TestClient, revision: dict[str, object]) -> dict[str, object]:
    response = client.post(
        f"/api/v1/boq/revisions/{revision['revision_id']}/issue",
        headers={**OWNER_HEADERS, "Idempotency-Key": _key("issue")},
        json={"expected_version": revision["version"]},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _accept(
    client: TestClient,
    revision: dict[str, object],
    *,
    replacement: dict[str, object] | None = None,
) -> dict[str, object]:
    payload: dict[str, object] = {
        "expected_version": revision["version"],
        "agreed_date": date(2026, 9, 17).isoformat(),
        "evidence_reference": "CRM-DEAL-123",
        "note": "บันทึกข้อตกลงโดยเจ้าของโครงการ",
    }
    if replacement is not None:
        payload["replacement"] = replacement
    response = client.post(
        f"/api/v1/boq/revisions/{revision['revision_id']}/record-acceptance",
        headers={**OWNER_HEADERS, "Idempotency-Key": _key("accept")},
        json=payload,
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_issue_accept_snapshot_history_and_owner_guards(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = _project(client, "lifecycle")
        draft = _create_saved_main(client, project_id)

        preview = client.post(
            f"/api/v1/boq/revisions/{draft['revision_id']}/preview-snapshots",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("preview")},
            json={"expected_version": draft["version"]},
        )
        assert preview.status_code == 201, preview.text
        frozen_preview = preview.json()["data"]
        assert frozen_preview["document"]["quotation"]["customer_name"].startswith(
            "บริษัท"
        )

        changed = client.patch(
            f"/api/v1/boq/revisions/{draft['revision_id']}",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("change")},
            json={
                "expected_version": draft["version"],
                "quotation": _quotation("บริษัท ลูกค้าใหม่ จำกัด"),
                "nodes": [
                    _scope_item(
                        row_id=draft["nodes"][0]["id"],
                        logical_id=draft["nodes"][0]["logical_id"],
                    )
                ],
            },
        )
        assert changed.status_code == 200, changed.text
        changed_revision = changed.json()["data"]
        pinned = client.get(
            f"/api/v1/boq/revisions/{draft['revision_id']}/preview",
            params={"snapshot_id": frozen_preview["snapshot"]["snapshot_id"]},
            headers=ADMIN_HEADERS,
        )
        assert pinned.status_code == 200, pinned.text
        assert pinned.json()["data"]["document"] == frozen_preview["document"]

        issued = _issue(client, changed_revision)
        assert issued["status"] == "ISSUED"
        assert issued["issued_snapshot_id"]
        workspace_before_accept = client.get(
            f"/api/v1/projects/{project_id}/boq-workspace", headers=ADMIN_HEADERS
        ).json()["data"]
        assert workspace_before_accept["active_source_kind"] == "LEGACY"
        assert workspace_before_accept["active_baseline_id"] is None

        immutable = client.patch(
            f"/api/v1/boq/revisions/{issued['revision_id']}",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("immutable")},
            json={"expected_version": issued["version"], "nodes": []},
        )
        assert immutable.status_code == 409
        assert immutable.json()["detail"]["code"] == "IMMUTABLE_BOQ_REVISION"

        admin_accept = client.post(
            f"/api/v1/boq/revisions/{issued['revision_id']}/record-acceptance",
            headers={**ADMIN_HEADERS, "Idempotency-Key": _key("admin-accept")},
            json={"expected_version": issued["version"], "agreed_date": "2026-09-17"},
        )
        assert admin_accept.status_code == 403

        accepted = _accept(client, issued)
        assert accepted["actor"].startswith("dev-admin@")
        assert accepted["evidence_reference"] == "CRM-DEAL-123"
        assert accepted["baseline"]["net_sell_ex_vat"] == "10000.00"
        accepted_revision = client.get(
            f"/api/v1/boq/revisions/{issued['revision_id']}", headers=ADMIN_HEADERS
        ).json()["data"]
        assert accepted_revision["status"] == "ACCEPTED"
        assert accepted_revision["accepted_agreed_date"] == "2026-09-17"

        revised = client.post(
            f"/api/v1/boq/revisions/{issued['revision_id']}/revise",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("revise")},
            json={"expected_version": accepted_revision["version"]},
        )
        assert revised.status_code == 201, revised.text
        revision_two = revised.json()["data"]
        assert revision_two["document_id"] == issued["document_id"]
        assert revision_two["document_number"] == issued["document_number"]
        assert revision_two["revision_number"] == 2
        assert revision_two["predecessor_revision_id"] == issued["revision_id"]

        wrong_snapshot = client.get(
            f"/api/v1/boq/revisions/{revision_two['revision_id']}/preview",
            params={"snapshot_id": issued["issued_snapshot_id"]},
            headers=ADMIN_HEADERS,
        )
        assert wrong_snapshot.status_code == 404
        client.portal.call(engine.dispose)


def test_alternative_and_change_orders_affect_only_accepted_baseline(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = _project(client, "baseline")
        logical_id = str(uuid4())
        main = _issue(
            client,
            _create_saved_main(
                client,
                project_id,
                amount="800.0000",
                logical_id=logical_id,
            ),
        )
        main_acceptance = _accept(client, main)
        main_baseline = main_acceptance["baseline"]

        alternative_response = client.post(
            f"/api/v1/boq/revisions/{main['revision_id']}/alternatives",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("alternative")},
            json={"expected_version": main_acceptance and main["version"] + 1},
        )
        assert alternative_response.status_code == 201, alternative_response.text
        alternative = alternative_response.json()["data"]
        alternative_save = client.patch(
            f"/api/v1/boq/revisions/{alternative['revision_id']}",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("alternative-save")},
            json={
                "expected_version": alternative["version"],
                "quotation": _quotation("บริษัท ทางเลือก จำกัด"),
                "nodes": [
                    _scope_item(
                        row_id=alternative["nodes"][0]["id"],
                        logical_id=logical_id,
                        material="1000.0000",
                        labor="200.0000",
                    )
                ],
            },
        )
        assert alternative_save.status_code == 200, alternative_save.text
        alternative_issued = _issue(client, alternative_save.json()["data"])

        missing_replacement = client.post(
            f"/api/v1/boq/revisions/{alternative['revision_id']}/record-acceptance",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("missing-replace")},
            json={
                "expected_version": alternative_issued["version"],
                "agreed_date": "2026-09-17",
            },
        )
        assert missing_replacement.status_code == 409
        assert (
            missing_replacement.json()["detail"]["code"]
            == "BASELINE_REPLACEMENT_CONFIRMATION_REQUIRED"
        )
        alternative_acceptance = _accept(
            client,
            alternative_issued,
            replacement={
                "baseline_id": main_baseline["baseline_id"],
                "baseline_version": main_baseline["baseline_version"],
                "retain_change_order_revision_ids": [],
                "absorb_change_order_revision_ids": [],
            },
        )
        assert alternative_acceptance["baseline"]["net_sell_ex_vat"] == "12000.00"
        assert alternative_acceptance["baseline"]["main_revision_id"] == alternative["revision_id"]

        add_created = client.post(
            f"/api/v1/projects/{project_id}/boq/change-orders",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("add")},
            json={
                "baseline_id": alternative_acceptance["baseline"]["baseline_id"],
                "baseline_version": alternative_acceptance["baseline"]["baseline_version"],
                "direction": "ADD",
                "deductions": [],
            },
        )
        assert add_created.status_code == 201, add_created.text
        add = add_created.json()["data"]
        add_saved = client.patch(
            f"/api/v1/boq/revisions/{add['revision_id']}",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("add-save")},
            json={
                "expected_version": add["version"],
                "quotation": _quotation("บริษัท ทางเลือก จำกัด"),
                "nodes": [
                    _scope_item(
                        description="งานเพิ่ม",
                        quantity="1.0000",
                        material="1000.0000",
                        labor="0.0000",
                    )
                ],
            },
        )
        assert add_saved.status_code == 200, add_saved.text
        add_issued = _issue(client, add_saved.json()["data"])

        before_add_accept = client.get(
            f"/api/v1/projects/{project_id}/boq-workspace", headers=ADMIN_HEADERS
        ).json()["data"]
        assert before_add_accept["active_baseline_version"] == 2
        add_acceptance = _accept(client, add_issued)
        assert add_acceptance["baseline"]["net_sell_ex_vat"] == "13000.00"
        assert add_acceptance["baseline"]["accepted_change_order_revision_ids"] == [
            add["revision_id"]
        ]

        duplicate_accept = client.post(
            f"/api/v1/boq/revisions/{add['revision_id']}/record-acceptance",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("double")},
            json={
                "expected_version": add_issued["version"],
                "agreed_date": "2026-09-17",
            },
        )
        assert duplicate_accept.status_code == 409

        deduct_created = client.post(
            f"/api/v1/projects/{project_id}/boq/change-orders",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("deduct")},
            json={
                "baseline_id": add_acceptance["baseline"]["baseline_id"],
                "baseline_version": add_acceptance["baseline"]["baseline_version"],
                "direction": "DEDUCT",
                "deductions": [
                    {
                        "target_revision_id": alternative["revision_id"],
                        "target_logical_id": logical_id,
                        "quantity": "0.5000",
                    }
                ],
            },
        )
        assert deduct_created.status_code == 201, deduct_created.text
        deduct = deduct_created.json()["data"]
        deduct_saved = client.patch(
            f"/api/v1/boq/revisions/{deduct['revision_id']}",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("deduct-save")},
            json={
                "expected_version": deduct["version"],
                "quotation": _quotation("บริษัท ทางเลือก จำกัด"),
                "nodes": [
                    {
                        **_scope_item(
                            row_id=deduct["nodes"][0]["id"],
                            logical_id=logical_id,
                            quantity="0.5000",
                            material="1000.0000",
                            labor="200.0000",
                        ),
                        "components": [],
                    }
                ],
            },
        )
        assert deduct_saved.status_code == 200, deduct_saved.text
        deduct_issued = _issue(client, deduct_saved.json()["data"])
        deduct_acceptance = _accept(client, deduct_issued)
        assert deduct_acceptance["baseline"]["net_sell_ex_vat"] == "12400.00"
        assert set(deduct_acceptance["baseline"]["accepted_change_order_revision_ids"]) == {
            add["revision_id"],
            deduct["revision_id"],
        }

        stale_deduct = client.post(
            f"/api/v1/projects/{project_id}/boq/change-orders",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("stale-deduct")},
            json={
                "baseline_id": add_acceptance["baseline"]["baseline_id"],
                "baseline_version": add_acceptance["baseline"]["baseline_version"],
                "direction": "DEDUCT",
                "deductions": [
                    {
                        "target_revision_id": alternative["revision_id"],
                        "target_logical_id": logical_id,
                        "quantity": "10.0000",
                    }
                ],
            },
        )
        assert stale_deduct.status_code == 409
        assert stale_deduct.json()["detail"]["code"] == "STALE_BASELINE_VERSION"

        over_deduct = client.post(
            f"/api/v1/projects/{project_id}/boq/change-orders",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("over-deduct")},
            json={
                "baseline_id": deduct_acceptance["baseline"]["baseline_id"],
                "baseline_version": deduct_acceptance["baseline"]["baseline_version"],
                "direction": "DEDUCT",
                "deductions": [
                    {
                        "target_revision_id": alternative["revision_id"],
                        "target_logical_id": logical_id,
                        "quantity": "10.0000",
                    }
                ],
            },
        )
        assert over_deduct.status_code == 409
        assert over_deduct.json()["detail"]["code"] == "DEDUCTION_EXCEEDS_REMAINING_SCOPE"

        workspace = client.get(
            f"/api/v1/projects/{project_id}/boq-workspace", headers=ADMIN_HEADERS
        ).json()["data"]
        assert workspace["active_source_kind"] == "V2"
        assert workspace["active_main_revision_id"] == alternative["revision_id"]
        assert len(workspace["active_change_order_revision_ids"]) == 2
        assert asyncio.run(_baseline_count(project_id)) == 1
        client.portal.call(engine.dispose)


def test_document_number_allocation_is_unique_under_concurrency(
) -> None:
    async def create_all() -> list[str]:
        engine = create_async_engine(os.environ["PHASE3_DATABASE_URL"])
        session_factory = async_sessionmaker(engine, expire_on_commit=False)

        async def allocate() -> str:
            async with session_factory() as session, session.begin():
                return await next_document_number(session, document_kind="MAIN")

        try:
            return await asyncio.gather(*[allocate() for _index in range(8)])
        finally:
            await engine.dispose()

    numbers = asyncio.run(create_all())
    assert len(numbers) == len(set(numbers)) == 8
    assert all(number.startswith("QT-2026-") for number in numbers)


def test_export_artifact_authorization_ownership_and_private_download(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    uploaded: dict[str, object] = {}
    upload_calls: list[str] = []

    async def fake_upload(**kwargs):
        uploaded.update(kwargs)
        upload_calls.append(kwargs["artifact_id"])
        return (
            "gs://private-test-bucket/boq_exports/"
            f"{kwargs['project_id']}/{kwargs['artifact_id']}"
        )

    async def fake_signed_url(*, storage_key: str, expires_in_minutes: int):
        assert storage_key.startswith("gs://private-test-bucket/boq_exports/")
        assert expires_in_minutes == 15
        return "https://storage.example.invalid/signed?ttl=900"

    monkeypatch.setattr(
        "app.services.boq_export_service.upload_boq_export_artifact", fake_upload
    )
    monkeypatch.setattr(
        "app.services.boq_export_service.generate_signed_url_for_storage_key",
        fake_signed_url,
    )
    with TestClient(app) as client:
        project_a = _project(client, "export-a")
        project_b = _project(client, "export-b")
        issued_a = _issue(client, _create_saved_main(client, project_a))
        draft_b = _create_saved_main(client, project_b)

        internal_pdf = client.post(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/exports",
            headers={**ADMIN_HEADERS, "Idempotency-Key": _key("internal-pdf")},
            json={
                "snapshot_id": issued_a["issued_snapshot_id"],
                "audience": "INTERNAL",
                "file_format": "PDF",
            },
        )
        assert internal_pdf.status_code == 422

        cross_project = client.post(
            f"/api/v1/boq/revisions/{draft_b['revision_id']}/exports",
            headers={**ADMIN_HEADERS, "Idempotency-Key": _key("cross-project")},
            json={
                "snapshot_id": issued_a["issued_snapshot_id"],
                "audience": "CUSTOMER",
                "file_format": "XLSX",
            },
        )
        assert cross_project.status_code == 404
        assert cross_project.json()["detail"]["code"] == "QUOTATION_SNAPSHOT_NOT_FOUND"

        export_key = _key("customer-export")
        exported = client.post(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/exports",
            headers={**ADMIN_HEADERS, "Idempotency-Key": export_key},
            json={
                "snapshot_id": issued_a["issued_snapshot_id"],
                "audience": "CUSTOMER",
                "file_format": "XLSX",
            },
        )
        assert exported.status_code == 201, exported.text
        artifact = exported.json()["data"]
        assert artifact["status"] == "READY"
        assert artifact["project_id"] == project_a
        assert artifact["revision_id"] == issued_a["revision_id"]
        assert artifact["snapshot_id"] == issued_a["issued_snapshot_id"]
        assert artifact["calculation_version"] == "boq-v2-calc-v1"
        assert artifact["sha256"] and artifact["size_bytes"] > 0
        original_artifact_sha256 = artifact["sha256"]
        original_artifact_bytes = uploaded["file_bytes"]
        assert (
            hashlib.sha256(original_artifact_bytes).hexdigest()
            == original_artifact_sha256
        )
        assert "download_url" not in uploaded
        assert uploaded["project_id"] == project_a

        sheet = load_workbook(BytesIO(uploaded["file_bytes"]), data_only=False).active
        values = " ".join(
            str(cell.value or "") for row in sheet.iter_rows() for cell in row
        )
        assert "ต้นทุนรวม" not in values
        assert "forecast_margin" not in values
        assert sheet["B12"].value.startswith("'=")

        revised = client.post(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/revise",
            headers={**OWNER_HEADERS, "Idempotency-Key": _key("export-revise")},
            json={"expected_version": issued_a["version"]},
        )
        assert revised.status_code == 201, revised.text
        historical_internal = client.post(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/exports",
            headers={
                **ADMIN_HEADERS,
                "Idempotency-Key": _key("historical-internal-export"),
            },
            json={
                "snapshot_id": issued_a["issued_snapshot_id"],
                "audience": "INTERNAL",
                "file_format": "XLSX",
            },
        )
        assert historical_internal.status_code == 201, historical_internal.text
        assert historical_internal.json()["data"]["cost_plan_version"] == 1

        repeated = client.post(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/exports",
            headers={**ADMIN_HEADERS, "Idempotency-Key": export_key},
            json={
                "snapshot_id": issued_a["issued_snapshot_id"],
                "audience": "CUSTOMER",
                "file_format": "XLSX",
            },
        )
        assert repeated.status_code == 201, repeated.text
        assert repeated.json()["data"]["artifact_id"] == artifact["artifact_id"]
        assert repeated.json()["data"]["sha256"] == original_artifact_sha256
        assert upload_calls.count(artifact["artifact_id"]) == 1
        assert len(upload_calls) == 2

        wrong_owner = client.get(
            f"/api/v1/boq/revisions/{draft_b['revision_id']}/exports/{artifact['artifact_id']}",
            headers=ADMIN_HEADERS,
        )
        assert wrong_owner.status_code == 404

        download = client.get(
            f"/api/v1/boq/revisions/{issued_a['revision_id']}/exports/{artifact['artifact_id']}/download",
            headers=ADMIN_HEADERS,
        )
        assert download.status_code == 200, download.text
        assert download.json()["data"]["download_url"].startswith(
            "https://storage.example.invalid/signed"
        )
        assert download.json()["data"]["download_expires_in_minutes"] == 15
        client.portal.call(engine.dispose)


def test_concurrent_alternative_acceptance_keeps_one_active_main(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = _project(client, "alternative-race")
        main_issued = _issue(client, _create_saved_main(client, project_id))
        main_acceptance = _accept(client, main_issued)
        accepted_main = client.get(
            f"/api/v1/boq/revisions/{main_issued['revision_id']}",
            headers=ADMIN_HEADERS,
        ).json()["data"]

        alternatives: list[dict[str, object]] = []
        for index in range(2):
            created = client.post(
                f"/api/v1/boq/revisions/{accepted_main['revision_id']}/alternatives",
                headers={
                    **OWNER_HEADERS,
                    "Idempotency-Key": _key(f"race-alt-{index}"),
                },
                json={"expected_version": accepted_main["version"]},
            )
            assert created.status_code == 201, created.text
            alternatives.append(_issue(client, created.json()["data"]))

        async def race() -> list[object]:
            local_engine = create_async_engine(os.environ["PHASE3_DATABASE_URL"])
            factory = async_sessionmaker(local_engine, expire_on_commit=False)

            async def accept_one(index: int) -> object:
                async with factory() as session:
                    try:
                        result = await record_acceptance(
                            session,
                            revision_id=UUID(alternatives[index]["revision_id"]),
                            request=BOQV2RecordAcceptanceRequest(
                                expected_version=alternatives[index]["version"],
                                agreed_date=date(2026, 9, 17),
                                replacement={
                                    "baseline_id": main_acceptance["baseline"]["baseline_id"],
                                    "baseline_version": main_acceptance["baseline"]["baseline_version"],
                                    "retain_change_order_revision_ids": [],
                                    "absorb_change_order_revision_ids": [],
                                },
                            ),
                            actor="race-owner@example.com",
                            idempotency_key=_key(f"race-accept-{index}"),
                        )
                        await session.commit()
                        return result
                    except Exception as error:
                        await session.rollback()
                        return error

            try:
                return await asyncio.gather(accept_one(0), accept_one(1))
            finally:
                await local_engine.dispose()

        outcomes = asyncio.run(race())
        successes = [item for item in outcomes if not isinstance(item, Exception)]
        failures = [item for item in outcomes if isinstance(item, Exception)]
        assert len(successes) == 1
        assert len(failures) == 1
        assert isinstance(failures[0], BOQDomainError)
        assert failures[0].code == "STALE_BASELINE_VERSION"
        assert asyncio.run(_baseline_count(project_id)) == 1
        workspace = client.get(
            f"/api/v1/projects/{project_id}/boq-workspace", headers=ADMIN_HEADERS
        ).json()["data"]
        assert workspace["active_main_revision_id"] in {
            item["revision_id"] for item in alternatives
        }
        client.portal.call(engine.dispose)


async def _baseline_count(project_id: str) -> int:
    database_url = os.environ["PHASE3_DATABASE_URL"].replace(
        "postgresql+asyncpg://", "postgresql://", 1
    )
    connection = await asyncpg.connect(database_url)
    try:
        return int(
            await connection.fetchval(
                "SELECT count(*) FROM boq_v2_project_baselines "
                "WHERE project_id = $1 AND is_active",
                UUID(project_id),
            )
        )
    finally:
        await connection.close()
