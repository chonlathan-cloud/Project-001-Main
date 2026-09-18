from __future__ import annotations

import asyncio
import os
import subprocess
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.core.database import engine
from app.models.boq import BOQItem
from app.schemas.boq_cost_schema import BOQV2CostPublishRequest
from app.services.boq_cost_service import publish_cost_plan
from app.services.boq_domain_service import BOQDomainError
from app.services.fund_service import FundDomainError, calculate_fund_summary, post_allocation
from main import app


pytestmark = pytest.mark.skipif(
    not os.getenv("PHASE4_DATABASE_URL"),
    reason="PHASE4_DATABASE_URL isolated PostgreSQL target is required",
)

OWNER = {"X-Debug-Admin-Role": "owner"}
ADMIN = {"X-Debug-Admin-Role": "admin"}


def key(prefix: str) -> str:
    return f"phase4-{prefix}-{uuid4()}"


def create_project(client: TestClient, label: str) -> str:
    response = client.post(
        "/api/v1/projects",
        headers=OWNER,
        json={
            "name": f"Phase 4 {label} {uuid4()}",
            "project_type": "COMMERCIAL",
            "overhead_percent": 0,
            "profit_percent": 0,
            "vat_percent": 7,
            "contingency_budget": 0,
            "status": "ACTIVE",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]["project_id"]


def quotation() -> dict[str, object]:
    return {
        "title": "Phase 4 quotation",
        "customer_name": "Phase 4 Customer",
        "quotation_date": "2026-09-17",
        "valid_until": "2026-10-17",
        "currency": "THB",
        "vat_rate": "7.0000",
        "discount_type": "NONE",
        "discount_value": "0.0000",
        "payment_schedule": [{"label": "Complete", "percentage": "100.0000"}],
        "commercial_terms": ["Internal fixture"],
        "document_pages": ["BOQ"],
    }


def scope_item(*, logical_id: str | None = None, quantity: str = "10.0000", unit: str = "m") -> dict[str, object]:
    return {
        "logical_id": logical_id or str(uuid4()),
        "parent_logical_id": None,
        "node_kind": "ITEM",
        "inclusion_state": "REQUIRED",
        "position": 0,
        "item_code": f"P4-{uuid4().hex[:8]}",
        "description": "Copper conduit installation",
        "specification": "Grade A",
        "quantity": quantity,
        "unit": unit,
        "sell_material_unit_rate": "1000.0000",
        "sell_labor_unit_rate": "500.0000",
        "components": [
            {
                "component_type": "MATERIAL",
                "quantity_basis": "INHERITED",
                "unit": unit,
                "specification": "Grade A",
                "cost_state": "PRICED",
                "unit_rate": "600.0000",
            },
            {
                "component_type": "LABOR",
                "quantity_basis": "INHERITED",
                "unit": unit,
                "specification": "Grade A",
                "cost_state": "PRICED",
                "unit_rate": "200.0000",
            },
        ],
    }


def create_saved_main(client: TestClient, project_id: str, *, item: dict[str, object] | None = None) -> dict[str, object]:
    created = client.post(
        f"/api/v1/projects/{project_id}/boq/documents",
        headers={**OWNER, "Idempotency-Key": key("create")},
        json={},
    )
    assert created.status_code == 201, created.text
    revision = created.json()["data"]
    saved = client.patch(
        f"/api/v1/boq/revisions/{revision['revision_id']}",
        headers={**OWNER, "Idempotency-Key": key("save")},
        json={
            "expected_version": revision["version"],
            "quotation": quotation(),
            "nodes": [item or scope_item()],
        },
    )
    assert saved.status_code == 200, saved.text
    return saved.json()["data"]


def issue_accept(client: TestClient, revision: dict[str, object]) -> dict[str, object]:
    issued = client.post(
        f"/api/v1/boq/revisions/{revision['revision_id']}/issue",
        headers={**OWNER, "Idempotency-Key": key("issue")},
        json={"expected_version": revision["version"]},
    )
    assert issued.status_code == 200, issued.text
    accepted = client.post(
        f"/api/v1/boq/revisions/{revision['revision_id']}/record-acceptance",
        headers={**OWNER, "Idempotency-Key": key("accept")},
        json={
            "expected_version": issued.json()["data"]["version"],
            "agreed_date": "2026-09-17",
            "evidence_reference": "P4-ACCEPT",
        },
    )
    assert accepted.status_code == 200, accepted.text
    return accepted.json()["data"]


def create_vendor(client: TestClient, project_id: str, name: str) -> dict[str, object]:
    response = client.post(
        f"/api/v1/projects/{project_id}/vendors",
        headers={**OWNER, "Idempotency-Key": key("vendor")},
        json={"display_name": name, "commercial_reference": f"REF-{name}"},
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def create_offer(
    client: TestClient,
    *,
    project_id: str,
    revision_id: str,
    vendor_id: str,
    component: dict[str, object],
    quantity: str = "10.0000",
    unit: str = "m",
    specification: str = "Grade A",
    rate: str = "450.0000",
    expired: bool = False,
) -> dict[str, object]:
    today = date.today()
    response = client.post(
        f"/api/v1/projects/{project_id}/boq/revisions/{revision_id}/vendor-offers",
        headers={**OWNER, "Idempotency-Key": key("offer")},
        json={
            "vendor_id": vendor_id,
            "quotation_reference": f"Q-{uuid4().hex[:8]}",
            "quotation_date": (today - timedelta(days=30) if expired else today).isoformat(),
            "valid_until": (today - timedelta(days=1) if expired else today + timedelta(days=30)).isoformat(),
            "currency": "THB",
            "tax_basis": "EXCLUSIVE_VAT",
            "tax_rate": "7.0000",
            "discount_type": "NONE",
            "discount_value": "0",
            "included_charges": "Delivery included",
            "charges_amount": "0",
            "evidence_storage_key": f"gs://phase4-private-evidence/{uuid4()}",
            "evidence_filename": "vendor-quotation.pdf",
            "evidence_content_type": "application/pdf",
            "evidence_size_bytes": 4096,
            "notes": "Manual commercial offer",
            "lines": [{
                "component_id": component["id"],
                "component_type": component["component_type"],
                "offered_quantity": quantity,
                "unit": unit,
                "specification": specification,
                "unit_rate": rate,
                "discount_amount": "0",
                "charges_amount": "0",
            }],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def test_vendor_compare_select_publish_and_cost_semantics(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = create_project(client, "vendor-cost")
        revision = create_saved_main(client, project_id)
        sell_before = revision["net_sell_ex_vat"]
        acceptance = issue_accept(client, revision)

        plan_response = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/working", headers=OWNER, json={}
        )
        assert plan_response.status_code == 200, plan_response.text
        plan = plan_response.json()["data"]
        material = next(row for row in plan["components"] if row["component_type"] == "MATERIAL")
        labor = next(row for row in plan["components"] if row["component_type"] == "LABOR")
        assert plan["original_estimated_cost"] == "8000.00"
        assert plan["estimated_cost"] == "8000.00"
        assert plan["agreed_cost"] is None
        assert plan["forecast_cost"] == "8000.00"

        vendor_a = create_vendor(client, project_id, "Material Partner")
        vendor_b = create_vendor(client, project_id, "Labor Partner")
        admin_vendor = client.post(
            f"/api/v1/projects/{project_id}/vendors",
            headers={**ADMIN, "Idempotency-Key": key("admin-vendor")},
            json={"display_name": "Admin cannot mutate"},
        )
        assert admin_vendor.status_code == 403

        partial = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor_a["id"], component=material, quantity="5.0000",
        )
        rejected = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("partial-select")},
            json={
                "offer_line_id": partial["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert rejected.status_code == 409, rejected.text
        assert rejected.json()["detail"]["code"] == "OFFER_FULL_COVERAGE_REQUIRED"

        spec_mismatch = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor_a["id"], component=material, specification="Grade B",
        )
        spec_rejected = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("spec-reject")},
            json={
                "offer_line_id": spec_mismatch["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert spec_rejected.status_code == 409
        assert spec_rejected.json()["detail"]["code"] == "OFFER_WARNING_CONFIRMATION_REQUIRED"

        material_offer = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor_a["id"], component=material, rate="450.0000",
        )
        labor_offer = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor_b["id"], component=labor, rate="150.0000", expired=True,
        )
        assert material_offer["evidence_filename"] == "vendor-quotation.pdf"
        admin_selection = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**ADMIN, "Idempotency-Key": key("admin-select")},
            json={
                "offer_line_id": material_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert admin_selection.status_code == 403

        funds_before = client.get(f"/api/v1/projects/{project_id}/funds/summary", headers=OWNER)
        assert funds_before.status_code == 200, funds_before.text
        funds_before_data = funds_before.json()["data"]

        stale_selection = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("stale-select")},
            json={
                "offer_line_id": material_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"] + 1,
            },
        )
        assert stale_selection.status_code == 409
        assert stale_selection.json()["detail"]["code"] == "STALE_COST_PLAN_VERSION"

        select_key = key("select-material")
        selected_material = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": select_key},
            json={
                "offer_line_id": material_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert selected_material.status_code == 200, selected_material.text
        repeated = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": select_key},
            json={
                "offer_line_id": material_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert repeated.status_code == 200, repeated.text
        assert repeated.json()["data"]["id"] == selected_material.json()["data"]["id"]
        conflicting_retry = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": select_key},
            json={
                "offer_line_id": material_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
                "acknowledge_warning": True,
                "reason": "Different request body for the same key",
            },
        )
        assert conflicting_retry.status_code == 409
        assert conflicting_retry.json()["detail"]["code"] == "IDEMPOTENCY_CONFLICT"

        plan = client.get(f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER).json()["data"]
        replacement_offer = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor_b["id"], component=material, rate="400.0000",
        )
        replacement = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("replace-material")},
            json={
                "offer_line_id": replacement_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert replacement.status_code == 200, replacement.text
        selection_history = client.get(
            f"/api/v1/projects/{project_id}/cost-selections", headers=OWNER
        ).json()["data"]
        assert next(row for row in selection_history if row["id"] == selected_material.json()["data"]["id"])["status"] == "REPLACED"

        plan = client.get(f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER).json()["data"]
        expired_without_reason = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("expired-reject")},
            json={
                "offer_line_id": labor_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert expired_without_reason.status_code == 409
        selected_labor = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("expired-accept")},
            json={
                "offer_line_id": labor_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
                "acknowledge_warning": True,
                "reason": "Vendor reconfirmed validity by phone",
            },
        )
        assert selected_labor.status_code == 200, selected_labor.text

        funds_after_offer = client.get(f"/api/v1/projects/{project_id}/funds/summary", headers=OWNER)
        assert funds_after_offer.json()["data"]["version"] == funds_before_data["version"]
        assert funds_after_offer.json()["data"]["projected_boq_margin"] == funds_before_data["projected_boq_margin"]

        plan = client.get(f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER).json()["data"]
        publish_key = key("publish")
        publish_payload = {
            "expected_version": plan["expected_version"],
            "expected_baseline_version": acceptance["baseline"]["baseline_version"],
            "reason": "Owner confirmed full supplier coverage",
        }
        admin_publish = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/publish",
            headers={**ADMIN, "Idempotency-Key": key("admin-publish")},
            json=publish_payload,
        )
        assert admin_publish.status_code == 403
        published = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": publish_key}, json=publish_payload,
        )
        assert published.status_code == 200, published.text
        published_data = published.json()["data"]
        assert published_data["cost_plan"]["original_estimated_cost"] == "8000.00"
        assert published_data["cost_plan"]["agreed_cost"] == "5500.00"
        assert published_data["cost_plan"]["forecast_cost"] == "5500.00"
        assert published_data["completeness_state"] == "COMPLETE"

        repeated_publish = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": publish_key}, json=publish_payload,
        )
        assert repeated_publish.status_code == 200, repeated_publish.text
        assert repeated_publish.json()["data"]["baseline_id"] == published_data["baseline_id"]

        funds_after_publish = client.get(f"/api/v1/projects/{project_id}/funds/summary", headers=OWNER)
        assert funds_after_publish.status_code == 200
        # Phase 0's consumer-specific compatibility projection remains public
        # until Phase 5, but the source fingerprint/version must change now.
        assert funds_after_publish.json()["data"]["projected_boq_margin"] == funds_before_data["projected_boq_margin"]
        assert funds_after_publish.json()["data"]["version"] != funds_before_data["version"]
        reloaded = client.get(f"/api/v1/boq/revisions/{revision['revision_id']}", headers=OWNER)
        assert reloaded.json()["data"]["net_sell_ex_vat"] == sell_before
        client.portal.call(engine.dispose)


def test_basis_change_stales_selection_and_requires_explicit_reconfirmation(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = create_project(client, "stale-selection")
        revision = create_saved_main(client, project_id)
        plan = client.get(f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER).json()["data"]
        component = next(row for row in revision["nodes"][0]["components"] if row["component_type"] == "MATERIAL")
        vendor = create_vendor(client, project_id, "Basis Vendor")
        offer = create_offer(
            client, project_id=project_id, revision_id=revision["revision_id"],
            vendor_id=vendor["id"], component=component,
        )
        selected = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("basis-select")},
            json={
                "offer_line_id": offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert selected.status_code == 200, selected.text

        changed_node = revision["nodes"][0]
        changed_node["unit"] = "lot"
        for row in changed_node["components"]:
            row["unit"] = "lot"
        changed_node.pop("depth", None)
        changed_node.pop("display_path", None)
        changed_node.pop("sell_total", None)
        for row in changed_node["components"]:
            row.pop("basis_version", None)
            row.pop("total", None)
            if row["quantity_basis"] == "INHERITED":
                row["quantity"] = None
        changed = client.patch(
            f"/api/v1/boq/revisions/{revision['revision_id']}",
            headers={**OWNER, "Idempotency-Key": key("basis-save")},
            json={"expected_version": revision["version"], "quotation": quotation(), "nodes": [changed_node]},
        )
        assert changed.status_code == 200, changed.text
        selections = client.get(f"/api/v1/projects/{project_id}/cost-selections", headers=OWNER)
        assert selections.status_code == 200
        current = next(row for row in selections.json()["data"] if row["current"])
        assert current["stale"] is True
        assert "STALE_COMPONENT_BASIS" in current["warnings"]

        current_plan = client.get(f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER).json()["data"]
        reconfirmed = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("basis-reconfirm")},
            json={
                "offer_line_id": offer["lines"][0]["id"],
                "expected_cost_plan_version": current_plan["expected_version"],
                "acknowledge_warning": True,
                "reason": "Vendor confirmed the same commercial basis despite unit label change",
            },
        )
        assert reconfirmed.status_code == 200, reconfirmed.text
        assert reconfirmed.json()["data"]["stale"] is False
        assert "UNIT_MISMATCH" in reconfirmed.json()["data"]["warnings"]
        client.portal.call(engine.dispose)


def test_unknown_zero_and_not_applicable_cost_states_remain_distinct(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        zero_project = create_project(client, "zero-na")
        zero_item = scope_item()
        zero_item["components"][0].update({
            "cost_state": "PRICED", "unit_rate": "0.0000",
            "explicit_zero_reason": "Vendor donated the material",
        })
        zero_item["components"][1].update({
            "cost_state": "NOT_APPLICABLE", "unit_rate": None,
        })
        zero_revision = create_saved_main(client, zero_project, item=zero_item)
        zero_acceptance = issue_accept(client, zero_revision)
        zero_plan = client.post(
            f"/api/v1/projects/{zero_project}/cost-plan/working", headers=OWNER, json={}
        ).json()["data"]
        assert zero_plan["completeness_state"] == "COMPLETE"
        assert zero_plan["original_estimated_cost"] == "0.00"
        assert zero_plan["estimated_cost"] == "0.00"
        assert zero_plan["forecast_cost"] == "0.00"
        zero_publish = client.post(
            f"/api/v1/projects/{zero_project}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": key("zero-publish")},
            json={
                "expected_version": zero_plan["expected_version"],
                "expected_baseline_version": zero_acceptance["baseline"]["baseline_version"],
                "reason": "Publish explicit zero and N/A fixture",
            },
        )
        assert zero_publish.status_code == 200, zero_publish.text
        assert zero_publish.json()["data"]["cost_plan"]["agreed_cost"] is None

        unknown_project = create_project(client, "unknown")
        unknown_item = scope_item()
        unknown_item["components"][0].update({"cost_state": "UNKNOWN", "unit_rate": None})
        unknown_item["components"][1].update({"cost_state": "NOT_APPLICABLE", "unit_rate": None})
        unknown_revision = create_saved_main(client, unknown_project, item=unknown_item)
        unknown_acceptance = issue_accept(client, unknown_revision)
        unknown_plan = client.post(
            f"/api/v1/projects/{unknown_project}/cost-plan/working", headers=OWNER, json={}
        ).json()["data"]
        assert unknown_plan["completeness_state"] == "INCOMPLETE"
        assert unknown_plan["original_estimated_cost"] is None
        assert unknown_plan["forecast_cost"] is None
        unknown_publish = client.post(
            f"/api/v1/projects/{unknown_project}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": key("unknown-publish")},
            json={
                "expected_version": unknown_plan["expected_version"],
                "expected_baseline_version": unknown_acceptance["baseline"]["baseline_version"],
                "reason": "Publish incomplete plan without inventing zero",
            },
        )
        assert unknown_publish.status_code == 200, unknown_publish.text
        assert unknown_publish.json()["data"]["completeness_state"] == "INCOMPLETE"
        assert unknown_publish.json()["data"]["cost_plan"]["forecast_cost"] is None
        client.portal.call(engine.dispose)


def test_catalog_promotion_history_reuse_and_permissions(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        source_project = create_project(client, "catalog-source")
        revision = create_saved_main(client, source_project)
        acceptance = issue_accept(client, revision)
        node = revision["nodes"][0]

        admin_promotion = client.post(
            "/api/v1/price-database/items/promote",
            headers=ADMIN,
            json={
                "source_node_id": node["id"], "code": node["item_code"],
                "name": node["description"], "unit": node["unit"],
                "reason": "Admin must not promote",
            },
        )
        assert admin_promotion.status_code == 403
        promoted = client.post(
            "/api/v1/price-database/items/promote",
            headers=OWNER,
            json={
                "source_node_id": node["id"], "code": node["item_code"],
                "name": node["description"], "specification": node["specification"],
                "unit": node["unit"], "category": "Electrical", "tags": ["conduit"],
                "price_kind": "MATERIAL_COST", "reference_amount": "600.0000",
                "tax_basis": "EXCLUSIVE_VAT", "effective_date": "2026-09-17",
                "reason": "Owner approved initial standard",
            },
        )
        assert promoted.status_code == 201, promoted.text
        item = promoted.json()["data"]

        history = client.get(f"/api/v1/price-database/items/{item['id']}", headers=ADMIN)
        assert history.status_code == 200, history.text
        kinds = {row["observation_kind"] for row in history.json()["data"]["observations"]}
        assert {"OFFERED_SELL", "ACCEPTED_SELL"}.issubset(kinds)
        assert history.json()["data"]["distinct_sample_count"] < len(history.json()["data"]["observations"])

        accepted_revision = client.get(
            f"/api/v1/boq/revisions/{revision['revision_id']}", headers=OWNER
        ).json()["data"]
        revised = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/revise",
            headers={**OWNER, "Idempotency-Key": key("catalog-revise")},
            json={"expected_version": accepted_revision["version"]},
        )
        assert revised.status_code == 201, revised.text
        revised_data = revised.json()["data"]
        revised_issue = client.post(
            f"/api/v1/boq/revisions/{revised_data['revision_id']}/issue",
            headers={**OWNER, "Idempotency-Key": key("catalog-reissue")},
            json={"expected_version": revised_data["version"]},
        )
        assert revised_issue.status_code == 200, revised_issue.text
        revised_history = client.get(
            f"/api/v1/price-database/items/{item['id']}", headers=ADMIN
        ).json()["data"]
        assert revised_history["distinct_sample_count"] == history.json()["data"]["distinct_sample_count"]
        assert len(revised_history["observations"]) > len(history.json()["data"]["observations"])

        updated = client.post(
            f"/api/v1/price-database/items/{item['id']}/reference-prices",
            headers=OWNER,
            json={
                "project_id": source_project,
                "expected_item_version": item["version"],
                "price_kind": "MATERIAL_COST", "amount": "625.0000",
                "tax_basis": "EXCLUSIVE_VAT", "effective_date": "2026-09-18",
                "reason": "Explicit reviewed reference update",
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["data"]["version"] == item["version"] + 1

        target_project = create_project(client, "catalog-reuse")
        target = create_saved_main(client, target_project, item=scope_item())
        reused = client.post(
            f"/api/v1/price-database/items/{item['id']}/reuse/{target['revision_id']}",
            headers={**OWNER, "Idempotency-Key": key("reuse")},
            json={
                "expected_revision_version": target["version"],
                "parent_logical_id": None, "position": 1,
                "quantity": "2.0000", "use_reference_prices": True,
            },
        )
        assert reused.status_code == 200, reused.text
        reused_node = next(row for row in reused.json()["data"]["nodes"] if row["catalog_item_id"] == item["id"])
        assert reused_node["catalog_item_version"] == item["version"] + 1
        assert reused_node["components"][0]["unit_rate"] == "625.0000"

        after_reuse = client.get(f"/api/v1/price-database/items/{item['id']}", headers=ADMIN)
        assert after_reuse.json()["data"]["item"]["version"] == item["version"] + 1

        source_plan = client.post(
            f"/api/v1/projects/{source_project}/cost-plan/working",
            headers=OWNER,
            json={},
        ).json()["data"]
        material = next(
            row
            for row in source_plan["components"]
            if row["component_type"] == "MATERIAL"
        )
        source_vendor = create_vendor(client, source_project, "Catalog Source Vendor")
        source_offer = create_offer(
            client,
            project_id=source_project,
            revision_id=revision["revision_id"],
            vendor_id=source_vendor["id"],
            component=material,
            rate="575.0000",
        )
        source_selection = client.post(
            f"/api/v1/projects/{source_project}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("catalog-selection")},
            json={
                "offer_line_id": source_offer["lines"][0]["id"],
                "expected_cost_plan_version": source_plan["expected_version"],
            },
        )
        assert source_selection.status_code == 200, source_selection.text
        source_plan = client.get(
            f"/api/v1/projects/{source_project}/cost-plan", headers=OWNER
        ).json()["data"]
        source_publish = client.post(
            f"/api/v1/projects/{source_project}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": key("catalog-publish")},
            json={
                "expected_version": source_plan["expected_version"],
                "expected_baseline_version": acceptance["baseline"]["baseline_version"],
                "reason": "Record catalog vendor provenance",
            },
        )
        assert source_publish.status_code == 200, source_publish.text

        vendor_history = client.get(
            f"/api/v1/price-database/items/{item['id']}", headers=ADMIN
        )
        vendor_kinds = {
            row["observation_kind"]
            for row in vendor_history.json()["data"]["observations"]
        }
        assert {
            "VENDOR_OFFERED_COST",
            "AGREED_VENDOR_COST",
            "ESTIMATED_COST",
        }.issubset(vendor_kinds)
        assert vendor_history.json()["data"]["item"]["version"] == item["version"] + 1

        archived = client.patch(
            f"/api/v1/price-database/items/{item['id']}",
            headers=OWNER,
            json={
                "project_id": source_project,
                "expected_version": item["version"] + 1,
                "name": item["name"],
                "specification": item["specification"],
                "unit": item["unit"],
                "category": item["category"],
                "tags": item["tags"],
                "status": "ARCHIVED",
                "reason": "Archive while retaining complete source history",
            },
        )
        assert archived.status_code == 200, archived.text
        archived_history = client.get(
            f"/api/v1/price-database/items/{item['id']}", headers=ADMIN
        )
        assert archived_history.json()["data"]["item"]["status"] == "ARCHIVED"
        assert len(archived_history.json()["data"]["observations"]) == len(
            vendor_history.json()["data"]["observations"]
        )

        cross_project_vendor = create_vendor(client, target_project, "Other Project Vendor")
        cross_offer = client.post(
            f"/api/v1/projects/{source_project}/boq/revisions/{revision['revision_id']}/vendor-offers",
            headers={**OWNER, "Idempotency-Key": key("cross")},
            json={
                "vendor_id": cross_project_vendor["id"], "quotation_date": "2026-09-17",
                "currency": "THB", "tax_basis": "NO_VAT", "discount_type": "NONE",
                "discount_value": "0", "lines": [{
                    "component_id": revision["nodes"][0]["components"][0]["id"],
                    "component_type": "MATERIAL", "offered_quantity": "10.0000",
                    "unit": "m", "specification": "Grade A", "unit_rate": "500.0000",
                }],
            },
        )
        assert cross_offer.status_code == 404
        client.portal.call(engine.dispose)


def test_cost_publish_serializes_with_funds_allocation(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        source_project = create_project(client, "publish-race-source")
        target_project = create_project(client, "publish-race-target")
        revision = create_saved_main(client, source_project)
        acceptance = issue_accept(client, revision)
        plan = client.post(
            f"/api/v1/projects/{source_project}/cost-plan/working", headers=OWNER, json={}
        ).json()["data"]

        async def race() -> tuple[object, object, int]:
            local_engine = create_async_engine(os.environ["PHASE4_DATABASE_URL"])
            factory = async_sessionmaker(local_engine, expire_on_commit=False)
            source_id = UUID(source_project)
            target_id = UUID(target_project)
            async with factory() as session:
                session.add_all([
                    BOQItem(
                        id=uuid4(), project_id=source_id, boq_type="CUSTOMER",
                        wbs_level=1, item_no="1", description="Legacy compatibility sell",
                        qty=1, grand_total=Decimal("15000.00"),
                    ),
                    BOQItem(
                        id=uuid4(), project_id=source_id, boq_type="SUBCONTRACTOR",
                        wbs_level=1, item_no="1", description="Legacy compatibility cost",
                        qty=1, grand_total=Decimal("8000.00"),
                    ),
                ])
                await session.commit()
            async with factory() as session:
                balance = await calculate_fund_summary(session, source_id)

            async def do_publish() -> object:
                async with factory() as session:
                    try:
                        result = await publish_cost_plan(
                            session,
                            project_id=source_id,
                            request=BOQV2CostPublishRequest(
                                expected_version=plan["expected_version"],
                                expected_baseline_version=acceptance["baseline"]["baseline_version"],
                                reason="Concurrent publish fixture",
                            ),
                            actor="race-owner@example.com",
                            idempotency_key=key("race-publish"),
                        )
                        await session.commit()
                        return result
                    except Exception as error:
                        await session.rollback()
                        return error

            async def do_allocate() -> object:
                async with factory() as session:
                    try:
                        return await post_allocation(
                            session,
                            source_project_id=source_id,
                            target_project_id=target_id,
                            amount=Decimal("100.00"),
                            currency="THB",
                            reason="Concurrent allocation fixture",
                            note=None,
                            expected_source_balance_version=balance.version,
                            idempotency_key=key("race-allocation"),
                            actor="race-owner@example.com",
                        )
                    except Exception as error:
                        return error

            publish_result, allocation_result = await asyncio.gather(do_publish(), do_allocate())
            async with factory() as session:
                refreshed = await calculate_fund_summary(session, source_id)
            await local_engine.dispose()
            return publish_result, allocation_result, int(refreshed.version != balance.version)

        publish_result, allocation_result, version_changed = asyncio.run(race())
        assert not isinstance(publish_result, Exception), publish_result
        assert version_changed == 1
        if isinstance(allocation_result, Exception):
            assert isinstance(allocation_result, FundDomainError)
            assert allocation_result.code == "STALE_FUND_BALANCE"
        client.portal.call(engine.dispose)


def test_concurrent_cost_publish_allows_one_winner(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    with TestClient(app) as client:
        project_id = create_project(client, "publish-race")
        revision = create_saved_main(client, project_id)
        acceptance = issue_accept(client, revision)
        plan = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/working", headers=OWNER, json={}
        ).json()["data"]

        async def race() -> list[object]:
            local_engine = create_async_engine(os.environ["PHASE4_DATABASE_URL"])
            factory = async_sessionmaker(local_engine, expire_on_commit=False)

            async def publish(index: int) -> object:
                async with factory() as session:
                    try:
                        result = await publish_cost_plan(
                            session,
                            project_id=UUID(project_id),
                            request=BOQV2CostPublishRequest(
                                expected_version=plan["expected_version"],
                                expected_baseline_version=acceptance["baseline"]["baseline_version"],
                                reason=f"Concurrent publish {index}",
                            ),
                            actor="race-owner@example.com",
                            idempotency_key=key(f"publish-{index}"),
                        )
                        await session.commit()
                        return result
                    except Exception as error:
                        await session.rollback()
                        return error

            try:
                return await asyncio.gather(publish(1), publish(2))
            finally:
                await local_engine.dispose()

        outcomes = asyncio.run(race())
        successes = [result for result in outcomes if not isinstance(result, Exception)]
        failures = [result for result in outcomes if isinstance(result, Exception)]
        assert len(successes) == 1
        assert len(failures) == 1
        assert isinstance(failures[0], BOQDomainError)
        assert failures[0].code in {"STALE_BASELINE_VERSION", "WORKING_COST_PLAN_REQUIRED"}
        client.portal.call(engine.dispose)


def test_rfq_and_selected_vendor_exports_are_snapshot_scoped_and_confidential(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(get_settings(), "boq_v2_enabled", True)
    uploaded: list[dict[str, object]] = []

    async def fake_upload(**kwargs):
        uploaded.append(kwargs)
        return f"gs://private-test-bucket/boq_exports/{kwargs['artifact_id']}"

    monkeypatch.setattr(
        "app.services.boq_export_service.upload_boq_export_artifact", fake_upload
    )
    with TestClient(app) as client:
        project_id = create_project(client, "vendor-export")
        revision = create_saved_main(client, project_id)
        issued_response = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/issue",
            headers={**OWNER, "Idempotency-Key": key("export-issue")},
            json={"expected_version": revision["version"]},
        )
        assert issued_response.status_code == 200, issued_response.text
        issued = issued_response.json()["data"]
        accepted_response = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/record-acceptance",
            headers={**OWNER, "Idempotency-Key": key("export-accept")},
            json={
                "expected_version": issued["version"],
                "agreed_date": "2026-09-17",
                "evidence_reference": "P4-EXPORT-ACCEPT",
            },
        )
        assert accepted_response.status_code == 200, accepted_response.text
        acceptance = accepted_response.json()["data"]

        vendor = create_vendor(client, project_id, "Chosen Vendor")
        competitor = create_vendor(client, project_id, "Competing Vendor")
        snapshot_component_id = revision["nodes"][0]["components"][0]["id"]

        rfq = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/exports",
            headers={**ADMIN, "Idempotency-Key": key("rfq-export")},
            json={
                "snapshot_id": issued["issued_snapshot_id"],
                "audience": "RFQ",
                "file_format": "XLSX",
                "vendor_id": vendor["id"],
                "selected_component_ids": [snapshot_component_id],
            },
        )
        assert rfq.status_code == 201, rfq.text
        assert rfq.json()["data"]["vendor_id"] == vendor["id"]
        rfq_sheet = load_workbook(BytesIO(uploaded[-1]["file_bytes"])).active
        assert rfq_sheet["G8"].value is None
        assert rfq_sheet["H8"].value is None
        rfq_values = " ".join(
            str(cell.value or "") for row in rfq_sheet.iter_rows() for cell in row
        )
        assert "Phase 4 Customer" not in rfq_values
        assert "1000.0000" not in rfq_values

        rfq_pdf = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/exports",
            headers={**ADMIN, "Idempotency-Key": key("rfq-pdf")},
            json={
                "snapshot_id": issued["issued_snapshot_id"],
                "audience": "RFQ",
                "file_format": "PDF",
                "vendor_id": vendor["id"],
                "selected_component_ids": [snapshot_component_id],
            },
        )
        assert rfq_pdf.status_code == 201, rfq_pdf.text
        assert uploaded[-1]["file_bytes"].startswith(b"%PDF-")
        rfq_pdf_path = tmp_path / "rfq.pdf"
        rfq_pdf_path.write_bytes(uploaded[-1]["file_bytes"])
        rfq_pdf_text = subprocess.run(
            ["pdftotext", str(rfq_pdf_path), "-"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        assert "Chosen Vendor" in rfq_pdf_text
        assert "1,000" not in rfq_pdf_text

        other_revision = create_saved_main(client, project_id)
        cross_revision = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/exports",
            headers={**ADMIN, "Idempotency-Key": key("rfq-cross-revision")},
            json={
                "snapshot_id": issued["issued_snapshot_id"],
                "audience": "RFQ",
                "file_format": "XLSX",
                "vendor_id": vendor["id"],
                "selected_component_ids": [
                    other_revision["nodes"][0]["components"][0]["id"]
                ],
            },
        )
        assert cross_revision.status_code == 422
        assert cross_revision.json()["detail"]["code"] == "EXPORT_SCOPE_INVALID"

        plan = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/working", headers=OWNER, json={}
        ).json()["data"]
        material = next(
            row for row in plan["components"] if row["component_type"] == "MATERIAL"
        )
        chosen_offer = create_offer(
            client,
            project_id=project_id,
            revision_id=revision["revision_id"],
            vendor_id=vendor["id"],
            component=material,
            rate="450.0000",
        )
        create_offer(
            client,
            project_id=project_id,
            revision_id=revision["revision_id"],
            vendor_id=competitor["id"],
            component=material,
            rate="777.0000",
        )
        selected = client.post(
            f"/api/v1/projects/{project_id}/cost-selections",
            headers={**OWNER, "Idempotency-Key": key("export-select")},
            json={
                "offer_line_id": chosen_offer["lines"][0]["id"],
                "expected_cost_plan_version": plan["expected_version"],
            },
        )
        assert selected.status_code == 200, selected.text
        plan = client.get(
            f"/api/v1/projects/{project_id}/cost-plan", headers=OWNER
        ).json()["data"]
        published = client.post(
            f"/api/v1/projects/{project_id}/cost-plan/publish",
            headers={**OWNER, "Idempotency-Key": key("export-publish")},
            json={
                "expected_version": plan["expected_version"],
                "expected_baseline_version": acceptance["baseline"]["baseline_version"],
                "reason": "Freeze selected vendor export fixture",
            },
        )
        assert published.status_code == 200, published.text
        cost_plan_version = published.json()["data"]["cost_plan"]["version"]

        vendor_export = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/exports",
            headers={**ADMIN, "Idempotency-Key": key("vendor-export")},
            json={
                "snapshot_id": issued["issued_snapshot_id"],
                "audience": "VENDOR",
                "file_format": "XLSX",
                "vendor_id": vendor["id"],
                "selected_component_ids": [snapshot_component_id],
                "cost_plan_version": cost_plan_version,
            },
        )
        assert vendor_export.status_code == 201, vendor_export.text
        artifact = vendor_export.json()["data"]
        assert artifact["vendor_id"] == vendor["id"]
        assert artifact["cost_plan_version"] == cost_plan_version
        vendor_sheet = load_workbook(BytesIO(uploaded[-1]["file_bytes"])).active
        assert vendor_sheet["G8"].value == 450
        assert vendor_sheet["H8"].value == 4500
        vendor_values = " ".join(
            str(cell.value or "") for row in vendor_sheet.iter_rows() for cell in row
        )
        assert "Competing Vendor" not in vendor_values
        assert "777" not in vendor_values
        assert "Phase 4 Customer" not in vendor_values
        assert "15000" not in vendor_values

        vendor_pdf = client.post(
            f"/api/v1/boq/revisions/{revision['revision_id']}/exports",
            headers={**ADMIN, "Idempotency-Key": key("vendor-pdf")},
            json={
                "snapshot_id": issued["issued_snapshot_id"],
                "audience": "VENDOR",
                "file_format": "PDF",
                "vendor_id": vendor["id"],
                "selected_component_ids": [snapshot_component_id],
                "cost_plan_version": cost_plan_version,
            },
        )
        assert vendor_pdf.status_code == 201, vendor_pdf.text
        vendor_pdf_path = tmp_path / "vendor.pdf"
        vendor_pdf_path.write_bytes(uploaded[-1]["file_bytes"])
        vendor_pdf_text = subprocess.run(
            ["pdftotext", str(vendor_pdf_path), "-"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
        assert "450.0000" in vendor_pdf_text
        assert "4,500.00" in vendor_pdf_text
        assert "777" not in vendor_pdf_text
        assert "Phase 4 Customer" not in vendor_pdf_text
        client.portal.call(engine.dispose)
